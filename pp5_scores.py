"""บันทึกคะแนนแบบ ปพ.5 (ตามหลักสูตรแกนกลางฯ 2551)

ขั้นตอนของครู
1. ตั้งค่าโครงสร้างวิชา: สัดส่วน เก็บ:กลางภาค:ปลายภาค (เช่น 50:20:30) และรายการตัวชี้วัด/ชิ้นงานของคะแนนเก็บ
2. บันทึกคะแนน: ตารางแบบ ปพ.5 (ตัวชี้วัด → รวมเก็บ → กลางภาค → ซ่อม → ปลายภาค → รวม → เกรด → หมายเหตุ)
   พร้อมร้อยละเวลาเรียนจากระบบเช็กชื่อ (ต่ำกว่า 80% = มส อัตโนมัติ)
3. คุณลักษณะอันพึงประสงค์ / อ่าน คิดวิเคราะห์ และเขียน: 4 ช่วง ช่วงละ 25 คะแนน → ระดับ 0-3
4. ส่งออกแบบคัดลอกไปวางในโปรแกรม SGS, พิมพ์ ปพ.5, ดูนักเรียนกลุ่มเสี่ยง (0 / ร / มส)

ข้อมูลแยกตามโรงเรียนและภาคเรียน ไม่ยุ่งกับสมุดคะแนนแบบเดิม (grades) ของระบบ

ใช้งาน: ``pp5_scores.register(app, globals())`` ใน app.py ก่อน init_db()
"""
import io
import math
import re
from datetime import datetime

from flask import render_template, request, redirect, url_for, flash, abort, send_file
from flask_login import login_required, current_user

GRADE_STEPS = [(80, '4'), (75, '3.5'), (70, '3'), (65, '2.5'), (60, '2'), (55, '1.5'), (50, '1'), (0, '0')]
GRADE_ORDER = ['4', '3.5', '3', '2.5', '2', '1.5', '1', '0', 'ร', 'มส']
TRAIT_LEVELS = [(80, 3, 'ดีเยี่ยม'), (70, 2, 'ดี'), (50, 1, 'ผ่าน'), (0, 0, 'ไม่ผ่าน')]
ATT_MIN = 80.0          # เวลาเรียนต่ำกว่านี้ = มส
SUBJECT_CODE = re.compile(r'^\s*([฀-๿]{1,2}\s?\d{5}[฀-๿]?)\s*[-–]?\s*(.*)$')
DEFAULT_ITEMS = [('ว1.1', 'ใบงาน/แบบฝึกหัด', 15), ('ว1.2', 'ชิ้นงานรายบุคคล', 15),
                 ('ว2.1', 'งานกลุ่ม/นำเสนอ', 10), ('ว2.2', 'ทดสอบย่อยระหว่างภาค', 10)]


def split_subject(name):
    m = SUBJECT_CODE.match(name or '')
    if m:
        return m.group(1).replace(' ', ''), (m.group(2) or '').strip() or name
    return '', name or ''


def grade_of(total):
    for lim, g in GRADE_STEPS:
        if total >= lim:
            return g
    return '0'


def trait_level(total):
    for lim, lv, _ in TRAIT_LEVELS:
        if total >= lim:
            return lv
    return 0


def fnum(v):
    """แปลงข้อความเป็นตัวเลข: ว่าง = None"""
    if v is None:
        return None
    s = str(v).strip().replace(',', '')
    if s == '' or s == '-':
        return None
    try:
        x = float(s)
    except ValueError:
        return None
    return x if math.isfinite(x) else None


def fmt(x):
    if x is None:
        return ''
    return str(int(x)) if float(x).is_integer() else ('%.2f' % x).rstrip('0').rstrip('.')


def register(app, g):
    db = g['db']
    User = g['User']
    Subject = g['Subject']
    Classroom = g['Classroom']
    ClassroomStudent = g['ClassroomStudent']
    Attendance = g['Attendance']
    role_required = g['role_required']
    get_active_semester = g['get_active_semester']
    owns_subject = g['owns_subject']
    owns_classroom = g['owns_classroom']
    deny_redirect = g['deny_redirect']
    pairs_for_user = g['phase1_subject_room_pairs']
    normalize_status = g['normalize_attendance_status']
    get_school_setting = g.get('get_school_setting')
    local_now = g.get('local_now') or datetime.now
    _current_school_id = g.get('current_school_id')

    def sid():
        try:
            return _current_school_id() if _current_school_id else getattr(current_user, 'school_id', None)
        except Exception:
            return getattr(current_user, 'school_id', None)

    # ------------------------------------------------------------------
    # Models
    # ------------------------------------------------------------------
    class P5Structure(db.Model):
        """โครงสร้างรายวิชาต่อภาคเรียน (ใช้กับทุกห้องของวิชานั้น)"""
        __tablename__ = 'p5_structure'
        id = db.Column(db.Integer, primary_key=True)
        school_id = db.Column(db.Integer, nullable=True, index=True)
        semester_id = db.Column(db.Integer, nullable=True, index=True)
        subject_id = db.Column(db.Integer, db.ForeignKey('subject.id'), nullable=False, index=True)
        collect_total = db.Column(db.Float, default=50)
        mid_total = db.Column(db.Float, default=20)
        final_total = db.Column(db.Float, default=30)
        mid_codes = db.Column(db.String(255), default='')
        mid_detail = db.Column(db.Text, default='')
        final_codes = db.Column(db.String(255), default='')
        final_detail = db.Column(db.Text, default='')
        total_periods = db.Column(db.Integer, nullable=True)   # ว่าง = ใช้จำนวนคาบของรายวิชา
        advisor1 = db.Column(db.String(150), default='')
        advisor2 = db.Column(db.String(150), default='')
        updated_by_id = db.Column(db.Integer, nullable=True)
        updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
        subject = db.relationship('Subject')

    class P5Item(db.Model):
        """ตัวชี้วัด/ชิ้นงานของคะแนนเก็บ"""
        __tablename__ = 'p5_item'
        id = db.Column(db.Integer, primary_key=True)
        structure_id = db.Column(db.Integer, db.ForeignKey('p5_structure.id'), nullable=False, index=True)
        seq = db.Column(db.Integer, default=1)
        code = db.Column(db.String(60), default='')
        name = db.Column(db.String(255), default='')
        detail = db.Column(db.Text, default='')
        max_score = db.Column(db.Float, default=10)
        half = db.Column(db.Integer, default=1)    # 1 = ก่อนกลางภาค, 2 = หลังกลางภาค (ใช้แยกช่อง SGS)

    class P5Score(db.Model):
        """คะแนนรายช่องของนักเรียน key: i<item_id> / mid / final / retake / remark / c1-c4 / r1-r4"""
        __tablename__ = 'p5_score'
        id = db.Column(db.Integer, primary_key=True)
        structure_id = db.Column(db.Integer, db.ForeignKey('p5_structure.id'), nullable=False, index=True)
        classroom_id = db.Column(db.Integer, nullable=True, index=True)
        student_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False, index=True)
        key = db.Column(db.String(20), nullable=False)
        value = db.Column(db.Float, nullable=True)
        text = db.Column(db.String(30), default='')
        updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)

    # ------------------------------------------------------------------
    # helpers
    # ------------------------------------------------------------------
    def get_structure(subject, create=False):
        sem = get_active_semester()
        sem_id = sem.id if sem else None
        st = P5Structure.query.filter_by(subject_id=subject.id, semester_id=sem_id).first()
        if not st and create:
            st = P5Structure(school_id=subject.school_id or sid(), semester_id=sem_id, subject_id=subject.id)
            db.session.add(st)
            db.session.flush()
            # คัดลอกโครงสร้างจากภาคเรียนก่อนหน้าของวิชาเดียวกัน ถ้ามี ไม่มีก็ใช้ค่าเริ่มต้น
            prev = P5Structure.query.filter(P5Structure.subject_id == subject.id, P5Structure.id != st.id) \
                .order_by(P5Structure.id.desc()).first()
            if prev:
                for f in ('collect_total', 'mid_total', 'final_total', 'mid_codes', 'mid_detail', 'final_codes',
                          'final_detail', 'total_periods', 'advisor1', 'advisor2'):
                    setattr(st, f, getattr(prev, f))
                for it in items_of(prev):
                    db.session.add(P5Item(structure_id=st.id, seq=it.seq, code=it.code, name=it.name,
                                          detail=it.detail, max_score=it.max_score, half=it.half))
            else:
                for i, (code, name, mx) in enumerate(DEFAULT_ITEMS, 1):
                    db.session.add(P5Item(structure_id=st.id, seq=i, code=code, name=name, max_score=mx,
                                          half=1 if i <= 2 else 2))
            db.session.commit()
        return st

    def items_of(st):
        return P5Item.query.filter_by(structure_id=st.id).order_by(P5Item.seq, P5Item.id).all() if st else []

    def students_of(room):
        links = ClassroomStudent.query.filter_by(classroom_id=room.id) \
            .join(User, ClassroomStudent.student_id == User.id).all()
        users = [l.student for l in links if l.student and getattr(l.student, 'is_active', True) is not False]

        def key(u):
            no = (getattr(u, 'student_no', '') or '').strip()
            return (0, int(no)) if no.isdigit() else (1, u.username or '', u.full_name or '')
        return sorted(users, key=key)

    def attendance_map(subject, room, sem, st):
        """ร้อยละเวลาเรียนแบบ ปพ.5: (คาบทั้งหมด − ขาด/ลา/โดด) ÷ คาบทั้งหมด"""
        q = Attendance.query.filter_by(subject_id=subject.id, classroom_id=room.id)
        if sem:
            q = q.filter(Attendance.date >= sem.start_date, Attendance.date <= sem.end_date)
        rows = q.all()
        total = int((st.total_periods if st and st.total_periods else None) or subject.total_periods or 40)
        taught = len({(r.date, r.period_no) for r in rows})
        out = {}
        for r in rows:
            s = normalize_status(r.status)
            d = out.setdefault(r.student_id, {'มา': 0, 'สาย': 0, 'ลา': 0, 'ขาด': 0, 'กิจกรรม': 0})
            if s == 'มา':
                d['มา'] += 1
            elif s == 'สาย':
                d['สาย'] += 1
            elif s == 'ลาป่วย':
                d['ลา'] += 1
            elif s in ('ขาด', 'โดดเรียน'):
                d['ขาด'] += 1
            elif s == 'ไปกิจกรรม':
                d['กิจกรรม'] += 1
        quota = int(total * (100 - ATT_MIN) / 100)
        for d in out.values():
            lost = d['ขาด'] + d['ลา']
            d['lost'] = lost
            d['pct'] = round(max(0.0, (total - lost) / total * 100), 2) if total else 100.0
            d['left'] = quota - lost
        return out, total, taught, quota

    def load_scores(st, student_ids):
        data = {}
        if not st or not student_ids:
            return data
        for s in P5Score.query.filter(P5Score.structure_id == st.id, P5Score.student_id.in_(student_ids)).all():
            data.setdefault(s.student_id, {})[s.key] = s
        return data

    def compute_row(st, items, sc, att):
        """คำนวณผลของนักเรียน 1 คน"""
        def val(k):
            r = sc.get(k)
            return r.value if r is not None else None
        item_vals = [val('i%d' % it.id) for it in items]
        collect = sum(min(v, it.max_score) for v, it in zip(item_vals, items) if v is not None)
        mid, final, retake = val('mid'), val('final'), val('retake')
        any_score = any(v is not None for v in item_vals + [mid, final])
        missing = [it.code or it.name for v, it in zip(item_vals, items) if v is None]
        if mid is None:
            missing.append('กลางภาค')
        if final is None:
            missing.append('ปลายภาค')
        total = collect + min(mid or 0, st.mid_total) + min(final or 0, st.final_total)
        total = round(total + 1e-9)
        remark_row = sc.get('remark')
        remark = (remark_row.text if remark_row else '') or ''
        pct = att['pct'] if att else None
        grade = grade_of(total) if any_score else ''
        auto_ms = pct is not None and pct < ATT_MIN
        if remark == 'มส' or auto_ms:
            grade = 'มส'
        elif remark == 'ร':
            grade = 'ร'
        elif grade == '0' and retake is not None:
            grade = '1' if retake >= 50 else '0'   # แก้ 0 ได้ไม่เกินระดับ 1
        traits = {}
        for kind in ('c', 'r'):
            vals = [val('%s%d' % (kind, p)) for p in range(1, 5)]
            tot = sum(v for v in vals if v is not None)
            traits[kind] = {'vals': vals, 'total': tot if any(v is not None for v in vals) else None,
                            'level': trait_level(tot) if any(v is not None for v in vals) else None}
        return {'items': item_vals, 'collect': collect, 'mid': mid, 'final': final, 'retake': retake,
                'total': total, 'grade': grade, 'remark': remark, 'auto_ms': auto_ms, 'missing': missing,
                'any': any_score, 'att': att, 'traits': traits}

    def build(subject, room):
        sem = get_active_semester()
        st = get_structure(subject, create=True)
        items = items_of(st)
        studs = students_of(room)
        scores = load_scores(st, [u.id for u in studs])
        att, total_periods, taught, quota = attendance_map(subject, room, sem, st)
        rows = []
        for i, u in enumerate(studs, 1):
            r = compute_row(st, items, scores.get(u.id, {}), att.get(u.id))
            r.update({'no': i, 'student': u, 'sno': (getattr(u, 'student_no', '') or '').strip() or u.username})
            rows.append(r)
        return {'sem': sem, 'st': st, 'items': items, 'rows': rows, 'total_periods': total_periods,
                'taught': taught, 'quota': quota}

    def check_pair(subject_id, classroom_id):
        subject = Subject.query.get_or_404(subject_id)
        room = Classroom.query.get_or_404(classroom_id)
        if not owns_subject(subject) or not owns_classroom(room):
            return None, None
        return subject, room

    def my_pairs():
        out = []
        for p in pairs_for_user():
            if p.subject and p.classroom and p.subject.is_active is not False and p.classroom.is_active is not False:
                out.append(p)
        return out

    def grade_dist(rows):
        dist = {g_: 0 for g_ in GRADE_ORDER}
        for r in rows:
            if r['grade'] in dist:
                dist[r['grade']] += 1
        return dist

    def trait_dist(rows, kind):
        dist = {3: 0, 2: 0, 1: 0, 0: 0}
        for r in rows:
            lv = r['traits'][kind]['level']
            if lv is not None:
                dist[lv] += 1
        return dist

    common = {'split_subject': split_subject, 'fmt': fmt, 'GRADE_ORDER': GRADE_ORDER, 'ATT_MIN': ATT_MIN}

    @app.context_processor
    def _pp5s_ctx():
        return {'pp5s_enabled': True}

    # ------------------------------------------------------------------
    # หน้าเลือกวิชา
    # ------------------------------------------------------------------
    @app.route('/pp5s')
    @login_required
    @role_required('teacher', 'admin')
    def pp5s_home():
        pairs = my_pairs()
        sel = request.args.get('pair', '')
        if sel and '-' in sel:
            a, b = sel.split('-', 1)
            return redirect(url_for('pp5s_entry', subject_id=int(a), classroom_id=int(b)))
        sem = get_active_semester()
        structs = {s.subject_id: s for s in P5Structure.query.filter_by(semester_id=sem.id if sem else None).all()}
        return render_template('pp5s/home.html', pairs=pairs, structs=structs, sem=sem, **common)

    # ------------------------------------------------------------------
    # ตั้งค่าโครงสร้างวิชา
    # ------------------------------------------------------------------
    @app.route('/pp5s/structure/<int:subject_id>', methods=['GET', 'POST'])
    @login_required
    @role_required('teacher', 'admin')
    def pp5s_structure(subject_id):
        subject = Subject.query.get_or_404(subject_id)
        if not owns_subject(subject):
            return deny_redirect('pp5s_home')
        st = get_structure(subject, create=True)
        back_room = request.values.get('room', type=int)
        if request.method == 'POST':
            for f in ('collect_total', 'mid_total', 'final_total'):
                v = fnum(request.form.get(f))
                if v is not None and v >= 0:
                    setattr(st, f, v)
            for f in ('mid_codes', 'mid_detail', 'final_codes', 'final_detail', 'advisor1', 'advisor2'):
                setattr(st, f, (request.form.get(f) or '').strip())
            tp = request.form.get('total_periods', '').strip()
            st.total_periods = int(tp) if tp.isdigit() and int(tp) > 0 else None
            st.updated_by_id = current_user.id
            ids = request.form.getlist('item_id')
            codes = request.form.getlist('item_code')
            names = request.form.getlist('item_name')
            details = request.form.getlist('item_detail')
            maxes = request.form.getlist('item_max')
            halves = request.form.getlist('item_half')
            existing = {it.id: it for it in items_of(st)}
            keep = set()
            seq = 0
            for i in range(len(codes)):
                code = (codes[i] or '').strip()
                name = (names[i] or '').strip() if i < len(names) else ''
                mx = fnum(maxes[i] if i < len(maxes) else None)
                if not code and not name:
                    continue
                seq += 1
                iid = int(ids[i]) if i < len(ids) and str(ids[i]).isdigit() else 0
                it = existing.get(iid)
                if not it:
                    it = P5Item(structure_id=st.id)
                    db.session.add(it)
                it.seq = seq
                it.code = code[:60]
                it.name = name[:255]
                it.detail = (details[i] if i < len(details) else '').strip()
                it.max_score = mx if mx is not None and mx >= 0 else 0
                it.half = 2 if (halves[i] if i < len(halves) else '1') == '2' else 1
                if iid:
                    keep.add(iid)
            for iid, it in existing.items():
                if iid not in keep:
                    P5Score.query.filter_by(structure_id=st.id, key='i%d' % iid).delete()
                    db.session.delete(it)
            db.session.commit()
            sum_items = sum(it.max_score for it in items_of(st))
            msg = 'บันทึกโครงสร้างวิชาแล้ว'
            if abs(sum_items - st.collect_total) > 0.001:
                flash(f'คะแนนเต็มของตัวชี้วัดรวม {fmt(sum_items)} ไม่เท่ากับคะแนนเก็บ {fmt(st.collect_total)} — ตรวจอีกครั้ง', 'warning')
            if abs(st.collect_total + st.mid_total + st.final_total - 100) > 0.001:
                flash(f'สัดส่วน เก็บ+กลางภาค+ปลายภาค = {fmt(st.collect_total + st.mid_total + st.final_total)} (ควรเป็น 100)', 'warning')
            flash(msg, 'success')
            if back_room:
                return redirect(url_for('pp5s_entry', subject_id=subject.id, classroom_id=back_room))
            return redirect(url_for('pp5s_structure', subject_id=subject.id))
        subjects = Subject.query.filter(Subject.id.in_({p.subject_id for p in my_pairs()} or [-1])).order_by(Subject.name).all()
        rooms = [p.classroom for p in my_pairs() if p.subject_id == subject.id]
        return render_template('pp5s/structure.html', subject=subject, st=st, items=items_of(st), subjects=subjects,
                               rooms=rooms, back_room=back_room, sem=get_active_semester(), **common)

    # ------------------------------------------------------------------
    # ตารางกรอกคะแนน
    # ------------------------------------------------------------------
    @app.route('/pp5s/<int:subject_id>/<int:classroom_id>', methods=['GET', 'POST'])
    @login_required
    @role_required('teacher', 'admin')
    def pp5s_entry(subject_id, classroom_id):
        subject, room = check_pair(subject_id, classroom_id)
        if not subject:
            return deny_redirect('pp5s_home')
        data = build(subject, room)
        st, items = data['st'], data['items']
        if request.method == 'POST':
            keys = ['i%d' % it.id for it in items] + ['mid', 'final', 'retake'] + \
                   ['c%d' % p for p in range(1, 5)] + ['r%d' % p for p in range(1, 5)]
            limits = {('i%d' % it.id): it.max_score for it in items}
            limits.update({'mid': st.mid_total, 'final': st.final_total, 'retake': 100})
            limits.update({('c%d' % p): 25 for p in range(1, 5)})
            limits.update({('r%d' % p): 25 for p in range(1, 5)})
            ids = [r['student'].id for r in data['rows']]
            existing = {}
            for s in P5Score.query.filter(P5Score.structure_id == st.id, P5Score.student_id.in_(ids or [-1])).all():
                existing[(s.student_id, s.key)] = s
            over = 0
            changed = 0
            for stu_id in ids:
                for k in keys + ['remark']:
                    field = 's_%d_%s' % (stu_id, k)
                    if field not in request.form:
                        continue
                    raw = request.form.get(field, '')
                    row = existing.get((stu_id, k))
                    if k == 'remark':
                        txt = raw if raw in ('ร', 'มส') else ''
                        if row is None and not txt:
                            continue
                        if row is None:
                            row = P5Score(structure_id=st.id, classroom_id=room.id, student_id=stu_id, key=k)
                            db.session.add(row)
                        if row.text != txt:
                            row.text = txt
                            changed += 1
                        continue
                    v = fnum(raw)
                    if v is not None:
                        v = max(0.0, v)
                        if v > limits.get(k, 100):
                            over += 1
                            v = limits.get(k, 100)
                    if row is None:
                        if v is None:
                            continue
                        row = P5Score(structure_id=st.id, classroom_id=room.id, student_id=stu_id, key=k)
                        db.session.add(row)
                    if row.value != v:
                        row.value = v
                        row.classroom_id = room.id
                        changed += 1
            db.session.commit()
            flash(f'บันทึกคะแนนแล้ว ({changed} ช่อง)', 'success')
            if over:
                flash(f'มี {over} ช่องที่เกินคะแนนเต็ม ระบบปรับเป็นคะแนนเต็มให้แล้ว', 'warning')
            return redirect(url_for('pp5s_entry', subject_id=subject.id, classroom_id=room.id,
                                    tab=request.form.get('tab', 'score')))
        pairs = my_pairs()
        return render_template('pp5s/entry.html', subject=subject, room=room, pairs=pairs,
                               tab=request.args.get('tab', 'score'), dist=grade_dist(data['rows']), **data, **common)

    # ------------------------------------------------------------------
    # พิมพ์ ปพ.5
    # ------------------------------------------------------------------
    @app.route('/pp5s/<int:subject_id>/<int:classroom_id>/print')
    @login_required
    @role_required('teacher', 'admin')
    def pp5s_print(subject_id, classroom_id):
        subject, room = check_pair(subject_id, classroom_id)
        if not subject:
            return deny_redirect('pp5s_home')
        data = build(subject, room)
        teacher = subject.teacher or current_user
        setting = get_school_setting() if get_school_setting else None
        return render_template('pp5s/print.html', subject=subject, room=room, teacher=teacher, setting=setting,
                               dist=grade_dist(data['rows']), cdist=trait_dist(data['rows'], 'c'),
                               rdist=trait_dist(data['rows'], 'r'), printed_at=local_now(), **data, **common)

    # ------------------------------------------------------------------
    # ส่งออก Excel แบบ ปพ.5
    # ------------------------------------------------------------------
    @app.route('/pp5s/<int:subject_id>/<int:classroom_id>/export.xlsx')
    @login_required
    @role_required('teacher', 'admin')
    def pp5s_export(subject_id, classroom_id):
        subject, room = check_pair(subject_id, classroom_id)
        if not subject:
            return deny_redirect('pp5s_home')
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
        data = build(subject, room)
        st, items = data['st'], data['items']
        wb = Workbook()
        ws = wb.active
        ws.title = 'ปพ.5'
        code, name = split_subject(subject.name)
        ws.append([f'ปพ.5 วิชา {code} {name} ห้อง {room.name} {data["sem"].name if data["sem"] else ""}'])
        head = ['ที่', 'รหัส', 'ชื่อ-สกุล', 'เวลาเรียน %'] + \
               [f'{it.code} {it.name} ({fmt(it.max_score)})' for it in items] + \
               [f'รวมเก็บ ({fmt(st.collect_total)})', f'กลางภาค ({fmt(st.mid_total)})', 'ซ่อม',
                f'ปลายภาค ({fmt(st.final_total)})', 'รวม (100)', 'เกรด', 'หมายเหตุ',
                'คุณลักษณะ รวม', 'คุณลักษณะ ระดับ', 'อ่านคิดฯ รวม', 'อ่านคิดฯ ระดับ']
        ws.append(head)
        for r in data['rows']:
            ws.append([r['no'], r['sno'], r['student'].full_name, r['att']['pct'] if r['att'] else None] +
                      r['items'] + [r['collect'], r['mid'], r['retake'], r['final'], r['total'],
                                    r['grade'], r['remark'] or ('มส' if r['auto_ms'] else ''),
                                    r['traits']['c']['total'], r['traits']['c']['level'],
                                    r['traits']['r']['total'], r['traits']['r']['level']])
        thin = Side(style='thin', color='999999')
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for c in ws[2]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='0F766E')
            c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
        ws['A1'].font = Font(bold=True, size=14)
        ws.column_dimensions['C'].width = 30
        ws.freeze_panes = 'D3'
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        fn = f'ปพ5_{code or subject.id}_{room.name.replace("/", "-")}.xlsx'
        return send_file(buf, as_attachment=True, download_name=fn,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    # ------------------------------------------------------------------
    # นักเรียนกลุ่มเสี่ยง (0 / ร / มส)
    # ------------------------------------------------------------------
    @app.route('/pp5s/risk')
    @login_required
    @role_required('teacher', 'admin')
    def pp5s_risk():
        groups = []
        counts = {'0': 0, 'ร': 0, 'มส': 0}
        for p in my_pairs():
            st = get_structure(p.subject, create=False)
            if not st:
                continue
            data = build(p.subject, p.classroom)
            bad = {'0': [], 'ร': [], 'มส': []}
            for r in data['rows']:
                if r['grade'] in bad:
                    bad[r['grade']].append(r)
            if any(bad.values()):
                for k in bad:
                    counts[k] += len(bad[k])
                groups.append({'pair': p, 'bad': bad})
        groups.sort(key=lambda x: (x['pair'].classroom.name, x['pair'].subject.name))
        return render_template('pp5s/risk.html', groups=groups, counts=counts, sem=get_active_semester(),
                               printing=request.args.get('print') == '1', **common)

    # ------------------------------------------------------------------
    # ส่งคะแนน ปพ.5 เข้าสมุดคะแนน BookMark (.accdb)
    # ต่อเข้าปุ่ม "ดึงจากระบบหลัก" ของหน้า /accdb/<sid> เดิม โดยไม่แก้โค้ด accdb ใน app.py
    #   UM01 = เก็บก่อนกลางภาค · UM02 = เก็บหลังกลางภาค · Midterm/Final = กลางภาค/ปลายภาค
    #   q1-q8 + ระดับ = คุณลักษณะฯ · l1-l5 + ระดับ = อ่าน คิดวิเคราะห์ เขียน · เกรด ร/มส/แก้ 0 ส่งไปด้วย
    # ------------------------------------------------------------------
    AccdbSubject = g.get('AccdbSubject')
    AccdbRow = g.get('AccdbRow')
    _orig_pull = app.view_functions.get('accdb_pull_preview')

    def _norm_name(*parts):
        x = ''.join((q or '') for q in parts)
        for pre in ('เด็กชาย', 'เด็กหญิง', 'นางสาว', 'นาย', 'นาง', 'ด.ช.', 'ด.ญ.'):
            x = x.replace(pre, '')
        return ''.join(x.split()).strip()

    def pp5_values_for_accdb(subject, room):
        """คืน (values ตามรหัส/ชื่อนักเรียน, weights) หรือ None ถ้าวิชานี้ยังไม่มีคะแนน ปพ.5"""
        sem = get_active_semester()
        st = P5Structure.query.filter_by(subject_id=subject.id, semester_id=sem.id if sem else None).first()
        if not st or not P5Score.query.filter_by(structure_id=st.id).first():
            return None
        data = build(subject, room)
        items = data['items']
        w1 = sum(it.max_score for it in items if it.half != 2)
        w2 = sum(it.max_score for it in items if it.half == 2)
        by_no, by_name = {}, {}
        for r in data['rows']:
            v = {}
            h1 = [min(x, it.max_score) for x, it in zip(r['items'], items) if x is not None and it.half != 2]
            h2 = [min(x, it.max_score) for x, it in zip(r['items'], items) if x is not None and it.half == 2]
            if h1:
                v['um01'] = round(sum(h1), 2)
            if h2:
                v['um02'] = round(sum(h2), 2)
            if r['mid'] is not None:
                v['mid'] = round(min(r['mid'], st.mid_total), 2)
            if r['final'] is not None:
                v['final'] = round(min(r['final'], st.final_total), 2)
            if r['any'] and r['grade'] in ('ร', 'มส') or (r['grade'] == '1' and r['retake'] is not None and r['total'] < 50):
                v['grade_override'] = r['grade']
            for kind, keys, gk in (('c', ['q%d' % i for i in range(1, 9)], 'qgrade'),
                                   ('r', ['l%d' % i for i in range(1, 6)], 'lgrade')):
                lv = r['traits'][kind]['level']
                if lv is not None:
                    for k in keys:
                        v[k] = lv
                    v[gk] = str(lv)
            u = r['student']
            no = (getattr(u, 'student_no', '') or '').strip()
            for key in filter(None, [no, (u.username or '').strip()]):
                by_no[key] = v
            nm = _norm_name(getattr(u, 'first_name', ''), getattr(u, 'last_name', '')) or _norm_name(u.full_name)
            if nm:
                by_name[nm] = v
        weights = {'um01': w1, 'mid': st.mid_total, 'um02': w2, 'final': st.final_total}
        return by_no, by_name, weights, len(data['rows'])

    if _orig_pull and AccdbSubject is not None and AccdbRow is not None:
        @login_required
        def accdb_pull_preview_pp5(sid):
            body = request.get_json(silent=True) or {}
            if body.get('source') == 'main':
                return _orig_pull(sid)
            sub = AccdbSubject.query.get(sid)
            if not sub or sub.owner_id != current_user.id:
                return _orig_pull(sid)
            subject = Subject.query.get(body.get('subject_id'))
            room = Classroom.query.get(body.get('classroom_id'))
            if not subject or not room or not owns_subject(subject):
                return _orig_pull(sid)
            got = pp5_values_for_accdb(subject, room)
            if not got:
                if body.get('source') == 'pp5':
                    return {'ok': False, 'error': 'วิชา/ห้องนี้ยังไม่มีคะแนนในระบบบันทึกคะแนน ปพ.5 ของภาคเรียนนี้'}
                return _orig_pull(sid)
            by_no, by_name, weights, n_main = got
            values, samples = {}, []
            arows = AccdbRow.query.filter_by(subject_id=sub.id).all()
            for ar in arows:
                key = (ar.sid or '').strip()
                v = by_no.get(key) or by_name.get(_norm_name(ar.first, ar.last))
                if v:
                    values[ar.sid] = v
                elif len(samples) < 6:
                    samples.append({'sid': ar.sid, 'name': ((ar.prefix or '') + (ar.first or '') + ' ' + (ar.last or '')).strip()})
            hint = (f'ดึงจาก ปพ.5: เก็บก่อนกลาง {fmt(weights["um01"])} · กลางภาค {fmt(weights["mid"])} · '
                    f'เก็บหลังกลาง {fmt(weights["um02"])} · ปลายภาค {fmt(weights["final"])} · พร้อมคุณลักษณะฯ/อ่านคิดวิเคราะห์ และเกรด ร/มส · '
                    '<b>ยังไม่บันทึก</b> จนกว่าจะกด “บันทึกลงฐานข้อมูล”')
            return {'ok': True, 'source': 'pp5', 'values': values, 'matched': len(values), 'weights': weights,
                    'hint': hint, 'debug': {'main_students': n_main, 'accdb_rows': len(arows), 'sample_accdb': samples,
                                            'sample_main': []}}

        app.view_functions['accdb_pull_preview'] = accdb_pull_preview_pp5

    return {'P5Structure': P5Structure, 'P5Item': P5Item, 'P5Score': P5Score}
