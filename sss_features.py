"""ฟีเจอร์ดูแลนักเรียน (ปรับแนวคิดมาจากระบบ SSS ของโรงเรียน แต่เขียนใหม่ทั้งหมด)

- [1] คะแนนความประพฤติ: เพิ่ม/ตัดรายคนหรือทั้งห้อง แนบหลักฐาน รายงานช่วงคะแนน และจัดอันดับ
- [2] เช็กชื่อกิจกรรมโฮมรูม พร้อมเรื่อง/หมายเหตุรายวัน
- [3] บัญชีของฉัน: เปลี่ยนรหัสผ่าน/ข้อมูลติดต่อ และบังคับเปลี่ยนรหัสเริ่มต้น
- [4] รายงานเช็กชื่อหน้าเสาธง/โฮมรูม/กิจกรรมโรงเรียน แบบตารางรายเดือน/รายเทอม และภาพรวมรายวัน
- [7] บันทึกอนุญาตออกนอกบริเวณโรงเรียน

ใช้งานโดยเรียก ``sss_features.register(app, globals())`` ใน app.py ก่อน ``init_db()``
เพื่อให้ db.create_all() สร้างตารางใหม่ให้อัตโนมัติ (ไม่แตะตารางเดิม)
"""
import os
from io import BytesIO
from datetime import datetime, date, timedelta

from flask import render_template, request, redirect, url_for, flash, abort, send_file
from flask_login import login_required, current_user


HR_STATUSES = ['มา', 'สาย', 'ขาด', 'ลาป่วย']          # สถานะที่ใช้เช็กโฮมรูม (ตรงกับระบบเดิม)
STATUS_LABEL = {'ลาป่วย': 'ลา', 'ไปกิจกรรม': 'กิจกรรม', 'โดดเรียน': 'โดด'}
ATTENDED_STATUSES = ('มา', 'สาย', 'ไปกิจกรรม')
REPORT_TYPES = {
    'assembly': 'เข้าแถวหน้าเสาธง',
    'homeroom': 'กิจกรรมโฮมรูม',
    'special': 'กิจกรรมโรงเรียน/กิจกรรมพิเศษ',
}
EVIDENCE_MAX_BYTES = 5 * 1024 * 1024

# รายการเริ่มต้น (แก้ไข/เพิ่ม/ปิดได้ที่หน้า ตั้งค่าคะแนนความประพฤติ)
DEFAULT_BEHAVIOR_ITEMS = [
    ('plus', 'ทำความสะอาดอาคารหรือบริเวณโรงเรียน', 10),
    ('plus', 'ให้ความร่วมมือเข้าร่วมกิจกรรมของโรงเรียน', 10),
    ('plus', 'เขตพื้นที่รับผิดชอบสะอาด', 10),
    ('plus', 'มีจิตสาธารณะ', 20),
    ('plus', 'เก็บของมีค่าแล้วส่งคืนเจ้าของ', 20),
    ('plus', 'ได้รับคัดเลือกเป็นตัวแทนโรงเรียนเข้าอบรมสัมมนา', 30),
    ('plus', 'ได้รับคัดเลือกเป็นตัวแทนโรงเรียนในการแสดงต่าง ๆ', 30),
    ('plus', 'ได้รับคัดเลือกเป็นตัวแทนแข่งขันทักษะวิชาการ', 50),
    ('plus', 'บริจาคโลหิต', 50),
    ('plus', 'ปฏิบัติหน้าที่พิเศษจากฝ่ายปกครอง', 50),
    ('minus', 'เหยียบส้นรองเท้า', 5),
    ('minus', 'ไม่ติดโบว์/ปล่อยสยายผม', 5),
    ('minus', 'นำขนมหรืออาหารมารับประทานบนอาคาร', 5),
    ('minus', 'ใส่รองเท้าขึ้นบนอาคารโดยไม่ได้รับอนุญาต', 5),
    ('minus', 'ไม่ทิ้งขยะลงถังหรือที่ทิ้งขยะ', 5),
    ('minus', 'แต่งกายไม่เรียบร้อย', 10),
    ('minus', 'ชายเสื้อออกนอกกางเกงหรือกระโปรง', 10),
    ('minus', 'กระโปรง/กางเกงสั้นหรือยาวเกินไป', 10),
    ('minus', 'ไม่รัดเข็มขัด/ใช้เข็มขัดผิดระเบียบ', 10),
    ('minus', 'ใส่เครื่องประดับผิดระเบียบมาโรงเรียน', 10),
    ('minus', 'ทรงผมผิดระเบียบ', 10),
    ('minus', 'สีผมไม่ถูกระเบียบ', 10),
    ('minus', 'เขตพื้นที่รับผิดชอบไม่สะอาด', 10),
    ('minus', 'ไม่ปฏิบัติตามกฎจราจร', 10),
    ('minus', 'หนีเรียนหรือขาดเรียนโดยไม่มีเหตุผล', 15),
    ('minus', 'ไม่ให้ความร่วมมือ สนับสนุนกิจกรรมของโรงเรียน', 15),
    ('minus', 'ซ่องสุมหรือก่อกวนความไม่สงบในหมู่คณะ', 20),
    ('minus', 'ออกนอกห้องเรียนโดยไม่ขออนุญาตครูผู้สอน', 20),
    ('minus', 'แสดงกิริยาไม่สุภาพต่อครู', 30),
    ('minus', 'ทำลายทรัพย์สินของโรงเรียน', 40),
    ('minus', 'ก่อการทะเลาะวิวาททั้งในและนอกโรงเรียน', 50),
    ('minus', 'เล่นการพนันทุกชนิด', 50),
    ('minus', 'ดื่มสุราหรือเกี่ยวข้องกับยาเสพติด', 50),
    ('minus', 'มีหรือเป็นผู้จำหน่ายยาเสพติด', 50),
    ('minus', 'ลักขโมยทรัพย์สินของผู้อื่น', 50),
    ('minus', 'ประพฤติตนให้โรงเรียนเสื่อมเสียชื่อเสียง', 50),
    ('minus', 'ข่มขู่เพื่อเรียกร้องหรือรีดไถทรัพย์สิน', 50),
    ('minus', 'ไม่ให้ความร่วมมือต่อกิจกรรมสำคัญของโรงเรียน', 50),
]
DEFAULT_BASE_SCORE = 100
DEFAULT_BANDS = '0,50,70,90,100'


def register(app, g):
    db = g['db']
    User = g['User']
    Classroom = g['Classroom']
    ClassroomStudent = g['ClassroomStudent']
    TeacherClassroom = g['TeacherClassroom']
    ClassroomActivity = g['ClassroomActivity']
    ActivityAttendance = g['ActivityAttendance']
    role_required = g['role_required']
    owns_classroom = g['owns_classroom']
    teacher_classroom_ids = g['teacher_classroom_ids']
    get_active_semester = g['get_active_semester']
    local_today = g['local_today']
    local_now = g['local_now']
    thai_date_label = g['thai_date_label']
    thai_date_short = g['thai_date_short']
    normalize_attendance_status = g['normalize_attendance_status']
    ATTENDANCE_SYMBOLS = g['ATTENDANCE_SYMBOLS']
    THAI_MONTHS = g['THAI_MONTHS']
    save_uploaded_file = g['save_uploaded_file']
    safe_remove_upload_file = g.get('safe_remove_upload_file')
    ensure_default_classroom_activities = g['ensure_default_classroom_activities']
    is_school_blocked_day = g['is_school_blocked_day']
    TEACHER_POSITION_CHOICES = g.get('TEACHER_POSITION_CHOICES', [])

    # ------------------------------------------------------------------
    # Models (ตารางใหม่ทั้งหมด ไม่แก้ตารางเดิม)
    # ------------------------------------------------------------------
    class SssConfig(db.Model):
        __tablename__ = 'sss_config'
        key = db.Column(db.String(80), primary_key=True)
        value = db.Column(db.Text, default='')

    class BehaviorItem(db.Model):
        __tablename__ = 'behavior_item'
        id = db.Column(db.Integer, primary_key=True)
        kind = db.Column(db.String(10), nullable=False, default='minus')  # plus / minus
        name = db.Column(db.String(255), nullable=False)
        points = db.Column(db.Integer, nullable=False, default=5)
        is_active = db.Column(db.Boolean, default=True)
        sort_order = db.Column(db.Integer, default=0)
        school_id = db.Column(db.Integer, nullable=True, index=True)

    class BehaviorRecord(db.Model):
        __tablename__ = 'behavior_record'
        id = db.Column(db.Integer, primary_key=True)
        student_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False, index=True)
        classroom_id = db.Column(db.Integer, db.ForeignKey('classroom.id'), nullable=True, index=True)
        item_id = db.Column(db.Integer, db.ForeignKey('behavior_item.id'), nullable=True)
        kind = db.Column(db.String(10), nullable=False, default='minus')
        item_name = db.Column(db.String(255), default='')
        points = db.Column(db.Integer, nullable=False, default=0)  # บวก=เพิ่ม ลบ=ตัด
        record_date = db.Column(db.Date, nullable=False)
        note = db.Column(db.Text, default='')
        evidence_path = db.Column(db.String(500), default='')
        evidence_name = db.Column(db.String(255), default='')
        recorded_by_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=True)
        semester_id = db.Column(db.Integer, db.ForeignKey('semester.id'), nullable=True, index=True)
        created_at = db.Column(db.DateTime, default=datetime.utcnow)
        school_id = db.Column(db.Integer, nullable=True, index=True)
        student = db.relationship('User', foreign_keys=[student_id])
        recorded_by = db.relationship('User', foreign_keys=[recorded_by_id])
        classroom = db.relationship('Classroom')

    class HomeroomLog(db.Model):
        __tablename__ = 'homeroom_log'
        id = db.Column(db.Integer, primary_key=True)
        classroom_id = db.Column(db.Integer, db.ForeignKey('classroom.id'), nullable=False, index=True)
        log_date = db.Column(db.Date, nullable=False)
        topic = db.Column(db.Text, default='')
        note = db.Column(db.Text, default='')
        recorded_by_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=True)
        updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
        school_id = db.Column(db.Integer, nullable=True, index=True)
        recorded_by = db.relationship('User')
        __table_args__ = (db.UniqueConstraint('classroom_id', 'log_date', name='uq_homeroom_room_date'),)

    class CampusLeave(db.Model):
        __tablename__ = 'campus_leave'
        id = db.Column(db.Integer, primary_key=True)
        student_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False, index=True)
        classroom_id = db.Column(db.Integer, db.ForeignKey('classroom.id'), nullable=True)
        destination = db.Column(db.String(255), default='')
        reason = db.Column(db.Text, default='')
        out_at = db.Column(db.DateTime, nullable=False)
        expect_return = db.Column(db.String(20), default='')
        returned_at = db.Column(db.DateTime, nullable=True)
        approved_by_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=True)
        returned_by_id = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=True)
        note = db.Column(db.Text, default='')
        semester_id = db.Column(db.Integer, db.ForeignKey('semester.id'), nullable=True)
        created_at = db.Column(db.DateTime, default=datetime.utcnow)
        school_id = db.Column(db.Integer, nullable=True, index=True)
        student = db.relationship('User', foreign_keys=[student_id])
        approved_by = db.relationship('User', foreign_keys=[approved_by_id])
        returned_by = db.relationship('User', foreign_keys=[returned_by_id])
        classroom = db.relationship('Classroom')

    for _cls in (SssConfig, BehaviorItem, BehaviorRecord, HomeroomLog, CampusLeave):
        g[_cls.__name__] = _cls

    # ------------------------------------------------------------------
    # Helpers — แยกข้อมูลตามโรงเรียน (multi-school)
    # ------------------------------------------------------------------
    _current_school_id = g.get('current_school_id')

    def sid():
        try:
            return _current_school_id() if _current_school_id else getattr(current_user, 'school_id', None)
        except Exception:
            return getattr(current_user, 'school_id', None)

    def scoped(model):
        """query ที่กรองเฉพาะโรงเรียนของผู้ใช้ปัจจุบัน"""
        if hasattr(model, 'school_id'):
            return model.query.filter(model.school_id == sid())
        return model.query

    def same_school(obj):
        return obj is not None and (not hasattr(obj, 'school_id') or getattr(obj, 'school_id', None) == sid())

    def school_student(student_id):
        return User.query.filter_by(id=student_id or -1, role='student').filter(User.school_id == sid()).first()

    _schema_checked = {'done': False}

    @app.before_request
    def _sss_ensure_schema():
        """เพิ่มคอลัมน์ school_id ให้ตารางของโมดูลนี้ ถ้าฐานข้อมูลเคยสร้างตารางไว้ก่อนมีคอลัมน์"""
        if _schema_checked['done']:
            return None
        _schema_checked['done'] = True
        try:
            from sqlalchemy import inspect as _insp, text as _text
            insp = _insp(db.engine)
            for table in ('behavior_item', 'behavior_record', 'homeroom_log', 'campus_leave'):
                if insp.has_table(table) and 'school_id' not in {c['name'] for c in insp.get_columns(table)}:
                    with db.engine.begin() as conn:
                        conn.execute(_text(f'ALTER TABLE {table} ADD COLUMN school_id INTEGER'))
        except Exception as e:  # อีก worker อาจเพิ่มไปแล้ว
            app.logger.warning('sss schema check: %s', e)
        return None

    def cfg_get(key, default=''):
        key = f's{sid()}:{key}'
        row = db.session.get(SssConfig, key)
        return row.value if row and row.value not in (None, '') else default

    def cfg_set(key, value):
        key = f's{sid()}:{key}'
        row = db.session.get(SssConfig, key)
        if not row:
            row = SssConfig(key=key)
            db.session.add(row)
        row.value = str(value)

    def base_score():
        try:
            return int(cfg_get('behavior_base_score', DEFAULT_BASE_SCORE))
        except ValueError:
            return DEFAULT_BASE_SCORE

    def band_edges():
        raw = cfg_get('behavior_bands', DEFAULT_BANDS)
        try:
            edges = sorted({int(x.strip()) for x in raw.split(',') if x.strip()})
        except ValueError:
            edges = [int(x) for x in DEFAULT_BANDS.split(',')]
        return edges or [0]

    def score_bands():
        """คืนรายการช่วงคะแนน [(label, lo, hi)] โดย lo/hi เป็น None ได้"""
        e = band_edges()
        bands = [(f'ต่ำกว่า {e[0]}', None, e[0] - 1)]
        for i in range(1, len(e)):
            lo = e[i - 1] if i == 1 else e[i - 1] + 1
            bands.append((f'{lo}–{e[i]}', lo, e[i]))
        bands.append((f'มากกว่า {e[-1]}', e[-1] + 1, None))
        return bands

    def band_of(score):
        for idx, (label, lo, hi) in enumerate(score_bands()):
            if (lo is None or score >= lo) and (hi is None or score <= hi):
                return idx, label
        return 0, ''

    def ensure_behavior_items():
        if scoped(BehaviorItem).count() == 0:
            for i, (kind, name, pts) in enumerate(DEFAULT_BEHAVIOR_ITEMS):
                db.session.add(BehaviorItem(kind=kind, name=name, points=pts, sort_order=i, is_active=True, school_id=sid()))
            db.session.commit()

    def active_rooms():
        return (scoped(Classroom)
                .filter(db.or_(Classroom.is_active == True, Classroom.is_active.is_(None)))
                .order_by(Classroom.name).all())

    def my_rooms():
        """ห้องที่ครูเกี่ยวข้อง (ประจำชั้น/ผูกห้อง/สอน) — แอดมินเห็นทุกห้อง"""
        if current_user.role == 'admin':
            return active_rooms()
        ids = set(teacher_classroom_ids() or [])
        return [r for r in active_rooms() if r.id in ids]

    def advisor_rooms():
        """ห้องที่ครูเป็นครูประจำชั้น/ครูที่ปรึกษา (ใช้กับโฮมรูม)"""
        if current_user.role == 'admin':
            return active_rooms()
        return [r for r in active_rooms() if owns_classroom(r)]

    def pick_room(rooms, room_id):
        if room_id:
            for r in rooms:
                if r.id == room_id:
                    return r
        return rooms[0] if rooms else None

    def room_links(room_id):
        return (ClassroomStudent.query
                .filter_by(classroom_id=room_id)
                .join(User, ClassroomStudent.student_id == User.id)
                .filter(db.or_(User.is_active == True, User.is_active.is_(None)))
                .order_by(User.username.asc(), User.full_name.asc()).all())

    def student_room(student_id):
        link = ClassroomStudent.query.filter_by(student_id=student_id).first()
        return link.classroom if link else None

    def parse_date(text, default=None):
        try:
            return datetime.strptime((text or '').strip(), '%Y-%m-%d').date()
        except ValueError:
            return default

    def sem_range(sem):
        if sem:
            return sem.start_date, sem.end_date
        today = local_today()
        return today.replace(month=1, day=1), today.replace(month=12, day=31)

    def search_students(q, limit=40):
        q = (q or '').strip()
        if not q:
            return []
        like = f'%{q}%'
        return (scoped(User).filter(User.role == 'student')
                .filter(db.or_(User.is_active == True, User.is_active.is_(None)))
                .filter(db.or_(User.username.ilike(like), User.full_name.ilike(like), User.student_no.ilike(like)))
                .order_by(User.username).limit(limit).all())

    def student_scores(student_ids, sem):
        """คืน {student_id: {'plus':x, 'minus':y, 'score':z, 'count':n}} ของภาคเรียน"""
        base = base_score()
        out = {sid: {'plus': 0, 'minus': 0, 'score': base, 'count': 0} for sid in student_ids}
        if not student_ids:
            return out
        q = scoped(BehaviorRecord).filter(BehaviorRecord.student_id.in_(list(student_ids)))
        q = q.filter(BehaviorRecord.semester_id == (sem.id if sem else None))
        for r in q.all():
            row = out.setdefault(r.student_id, {'plus': 0, 'minus': 0, 'score': base, 'count': 0})
            if r.points >= 0:
                row['plus'] += r.points
            else:
                row['minus'] += -r.points
            row['score'] += r.points
            row['count'] += 1
        return out

    def file_size(f):
        try:
            pos = f.stream.tell()
            f.stream.seek(0, os.SEEK_END)
            size = f.stream.tell()
            f.stream.seek(pos)
            return size
        except Exception:
            return 0

    def can_manage_record(rec):
        return same_school(rec) and (current_user.role == 'admin' or rec.recorded_by_id == current_user.id)

    def xlsx_response(wb, filename):
        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)
        return send_file(bio, as_attachment=True, download_name=filename,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def status_label(st):
        return STATUS_LABEL.get(st, st)

    @app.context_processor
    def _sss_ctx():
        return {'sss_status_label': status_label}

    # ------------------------------------------------------------------
    # [3] บัญชีของฉัน + บังคับเปลี่ยนรหัสผ่านเริ่มต้น
    # ------------------------------------------------------------------
    ALLOWED_WHILE_MUST_CHANGE = {'account', 'logout', 'static', 'login', 'uploaded_file'}

    @app.before_request
    def _force_password_change():
        if os.environ.get('FORCE_PASSWORD_CHANGE', '1') == '0':
            return None
        try:
            if not current_user.is_authenticated or not getattr(current_user, 'must_change_password', False):
                return None
        except Exception:
            return None
        ep = request.endpoint or ''
        if ep in ALLOWED_WHILE_MUST_CHANGE or ep.startswith('static'):
            return None
        if request.method == 'GET':
            flash('กรุณาตั้งรหัสผ่านใหม่ก่อนใช้งาน เพื่อความปลอดภัยของข้อมูล', 'warning')
        return redirect(url_for('account'))

    def password_problem(user, new, confirm, current):
        if len(new) < 6:
            return 'รหัสผ่านใหม่ต้องยาวอย่างน้อย 6 ตัวอักษร'
        if new != confirm:
            return 'รหัสผ่านใหม่ทั้งสองช่องไม่ตรงกัน'
        if new == current:
            return 'รหัสผ่านใหม่ต้องไม่ซ้ำรหัสเดิม'
        if new == (user.username or '') or new == (user.citizen_id or ''):
            return 'ห้ามใช้ชื่อผู้ใช้หรือเลขบัตรประชาชนเป็นรหัสผ่าน'
        bd = ''.join(ch for ch in (user.birth_date or '') if ch.isdigit())
        if bd and ''.join(ch for ch in new if ch.isdigit()) == bd and new.isdigit():
            return 'ห้ามใช้วันเกิดเป็นรหัสผ่าน'
        if new in ('123456', '1234567', '12345678', '111111', '000000', 'password'):
            return 'รหัสผ่านนี้เดาง่ายเกินไป'
        return ''

    @app.route('/account', methods=['GET', 'POST'])
    @login_required
    def account():
        user = current_user
        if request.method == 'POST':
            action = request.form.get('action')
            if action == 'password':
                current = request.form.get('current_password', '')
                new = request.form.get('new_password', '')
                confirm = request.form.get('confirm_password', '')
                if not user.check_password(current):
                    flash('รหัสผ่านปัจจุบันไม่ถูกต้อง', 'danger')
                else:
                    problem = password_problem(user, new, confirm, current)
                    if problem:
                        flash(problem, 'danger')
                    else:
                        user.set_password(new)
                        user.must_change_password = False
                        db.session.commit()
                        flash('เปลี่ยนรหัสผ่านเรียบร้อยแล้ว', 'success')
                        return redirect(url_for('index'))
            elif action == 'profile':
                user.phone = (request.form.get('phone') or '').strip()[:50]
                if user.role == 'student':
                    user.guardian_name = (request.form.get('guardian_name') or '').strip()[:255]
                    user.guardian_phone = (request.form.get('guardian_phone') or '').strip()[:50]
                else:
                    name = (request.form.get('full_name') or '').strip()
                    if name:
                        user.full_name = name[:255]
                    pos = (request.form.get('position') or '').strip()
                    if pos:
                        user.position = pos[:80]
                db.session.commit()
                flash('บันทึกข้อมูลส่วนตัวแล้ว', 'success')
            return redirect(url_for('account'))
        return render_template('sss/account.html', user=user, position_choices=TEACHER_POSITION_CHOICES,
                               must_change=getattr(user, 'must_change_password', False))

    # ------------------------------------------------------------------
    # [1] คะแนนความประพฤติ
    # ------------------------------------------------------------------
    @app.route('/behavior')
    @login_required
    @role_required('teacher', 'admin')
    def behavior_home():
        ensure_behavior_items()
        sem = get_active_semester()
        rooms = active_rooms()
        mine = advisor_rooms()
        room = pick_room(rooms, request.args.get('room_id', type=int) or (mine[0].id if mine else None))
        rows = []
        if room:
            links = room_links(room.id)
            scores = student_scores([l.student_id for l in links], sem)
            for l in links:
                s = scores.get(l.student_id)
                _, band = band_of(s['score'])
                rows.append({'student': l.student, **s, 'band': band})
        today = local_today()
        recent = (scoped(BehaviorRecord).filter(BehaviorRecord.semester_id == (sem.id if sem else None))
                  .order_by(BehaviorRecord.created_at.desc()).limit(8).all())
        today_count = scoped(BehaviorRecord).filter(BehaviorRecord.record_date == today).count()
        return render_template('sss/behavior_home.html', rooms=rooms, room=room, rows=rows, sem=sem,
                               base=base_score(), recent=recent, today_count=today_count)

    @app.route('/behavior/record', methods=['GET', 'POST'])
    @login_required
    @role_required('teacher', 'admin')
    def behavior_record():
        ensure_behavior_items()
        sem = get_active_semester()
        kind = request.values.get('kind', 'minus')
        if kind not in ('plus', 'minus'):
            kind = 'minus'
        if request.method == 'POST':
            item = db.session.get(BehaviorItem, request.form.get('item_id', type=int) or -1)
            if item is not None and not same_school(item):
                item = None
            student_ids = [int(x) for x in request.form.getlist('student_ids') if str(x).isdigit()]
            rec_date = parse_date(request.form.get('record_date'), local_today())
            note = (request.form.get('note') or '').strip()
            back = url_for('behavior_record', kind=kind, room_id=request.form.get('room_id') or None,
                           q=request.form.get('q') or None)
            if not item or not item.is_active:
                flash('กรุณาเลือกรายการความประพฤติ', 'danger')
                return redirect(back)
            if not student_ids:
                flash('กรุณาเลือกนักเรียนอย่างน้อย 1 คน', 'danger')
                return redirect(back)
            evidence_path, evidence_name = '', ''
            f = request.files.get('evidence')
            if f and f.filename:
                if file_size(f) > EVIDENCE_MAX_BYTES:
                    flash('ไฟล์หลักฐานต้องมีขนาดไม่เกิน 5 MB', 'danger')
                    return redirect(back)
                try:
                    evidence_path, evidence_name = save_uploaded_file(f, subdir='behavior')
                except ValueError as e:
                    flash(str(e), 'danger')
                    return redirect(back)
            signed = item.points if item.kind == 'plus' else -abs(item.points)
            students = scoped(User).filter(User.id.in_(student_ids), User.role == 'student').all()
            for st in students:
                room = student_room(st.id)
                db.session.add(BehaviorRecord(
                    student_id=st.id, classroom_id=room.id if room else None, item_id=item.id,
                    kind=item.kind, item_name=item.name, points=signed, record_date=rec_date, note=note,
                    evidence_path=evidence_path, evidence_name=evidence_name,
                    recorded_by_id=current_user.id, semester_id=sem.id if sem else None, school_id=sid()))
            db.session.commit()
            word = 'เพิ่ม' if item.kind == 'plus' else 'ตัด'
            flash(f'{word}คะแนน "{item.name}" ({abs(signed)} คะแนน) ให้นักเรียน {len(students)} คนแล้ว', 'success')
            return redirect(back)

        rooms = active_rooms()
        room_id = request.args.get('room_id', type=int)
        q = (request.args.get('q') or '').strip()
        room = pick_room(rooms, room_id) if room_id else None
        if q:
            students = search_students(q)
            source = f'ผลค้นหา "{q}"'
        elif room:
            students = [l.student for l in room_links(room.id)]
            source = f'ห้อง {room.name}'
        else:
            students, source = [], ''
        scores = student_scores([s.id for s in students], sem)
        items = (scoped(BehaviorItem).filter_by(kind=kind, is_active=True)
                 .order_by(BehaviorItem.points.asc(), BehaviorItem.sort_order.asc()).all())
        rooms_of = {}
        if q and students:
            for link in ClassroomStudent.query.filter(ClassroomStudent.student_id.in_([s.id for s in students])).all():
                rooms_of.setdefault(link.student_id, link.classroom.name if link.classroom else '')
        return render_template('sss/behavior_record.html', rooms=rooms, room=room, q=q, kind=kind,
                               items=items, students=students, scores=scores, source=source,
                               rooms_of=rooms_of, today=local_today().isoformat(), sem=sem)

    @app.route('/behavior/records')
    @login_required
    @role_required('teacher', 'admin')
    def behavior_records():
        sem = get_active_semester()
        rooms = active_rooms()
        room_id = request.args.get('room_id', type=int)
        kind = request.args.get('kind', '')
        q = (request.args.get('q') or '').strip()
        d_from = parse_date(request.args.get('date_from'))
        d_to = parse_date(request.args.get('date_to'))
        mine_only = request.args.get('mine') == '1'
        query = scoped(BehaviorRecord).filter(BehaviorRecord.semester_id == (sem.id if sem else None))
        if room_id:
            query = query.filter(BehaviorRecord.classroom_id == room_id)
        if kind in ('plus', 'minus'):
            query = query.filter(BehaviorRecord.kind == kind)
        if d_from:
            query = query.filter(BehaviorRecord.record_date >= d_from)
        if d_to:
            query = query.filter(BehaviorRecord.record_date <= d_to)
        if mine_only:
            query = query.filter(BehaviorRecord.recorded_by_id == current_user.id)
        if q:
            like = f'%{q}%'
            query = (query.join(User, BehaviorRecord.student_id == User.id)
                     .filter(db.or_(User.full_name.ilike(like), User.username.ilike(like),
                                    BehaviorRecord.item_name.ilike(like))))
        records = query.order_by(BehaviorRecord.record_date.desc(), BehaviorRecord.id.desc()).limit(500).all()
        return render_template('sss/behavior_records.html', records=records, rooms=rooms, room_id=room_id,
                               kind=kind, q=q, d_from=d_from, d_to=d_to, mine_only=mine_only,
                               can_manage=can_manage_record, sem=sem)

    @app.route('/behavior/records/delete', methods=['POST'])
    @login_required
    @role_required('teacher', 'admin')
    def behavior_records_delete():
        ids = [int(x) for x in request.form.getlist('record_ids') if str(x).isdigit()]
        recs = scoped(BehaviorRecord).filter(BehaviorRecord.id.in_(ids or [-1])).all()
        deleted, skipped = 0, 0
        for r in recs:
            if not can_manage_record(r):
                skipped += 1
                continue
            path = r.evidence_path
            db.session.delete(r)
            deleted += 1
            if path and safe_remove_upload_file and not BehaviorRecord.query.filter(
                    BehaviorRecord.evidence_path == path, BehaviorRecord.id != r.id).first():
                try:
                    safe_remove_upload_file(path)
                except Exception:
                    pass
        db.session.commit()
        msg = f'ลบรายการแล้ว {deleted} รายการ'
        if skipped:
            msg += f' (ข้าม {skipped} รายการที่บันทึกโดยครูท่านอื่น)'
        flash(msg, 'success' if deleted else 'warning')
        return redirect(request.referrer or url_for('behavior_records'))

    def _student_history(student):
        sem = get_active_semester()
        recs = (scoped(BehaviorRecord).filter_by(student_id=student.id)
                .filter(BehaviorRecord.semester_id == (sem.id if sem else None))
                .order_by(BehaviorRecord.record_date.desc(), BehaviorRecord.id.desc()).all())
        s = student_scores([student.id], sem)[student.id]
        _, band = band_of(s['score'])
        return render_template('sss/behavior_student.html', student=student, records=recs, s=s, band=band,
                               base=base_score(), sem=sem, room=student_room(student.id),
                               can_manage=can_manage_record)

    @app.route('/behavior/student/<int:student_id>')
    @login_required
    @role_required('teacher', 'admin')
    def behavior_student(student_id):
        student = school_student(student_id)
        if not student:
            abort(404)
        return _student_history(student)

    @app.route('/behavior/me')
    @login_required
    def behavior_me():
        if current_user.role != 'student':
            return redirect(url_for('behavior_home'))
        return _student_history(current_user)

    def _report_data(room_scope):
        sem = get_active_semester()
        if room_scope == 'all':
            rooms = active_rooms()
        else:
            r = db.session.get(Classroom, room_scope)
            rooms = [r] if r and same_school(r) else []
        links = []
        for r in rooms:
            links.extend(room_links(r.id))
        seen, uniq = set(), []
        for l in links:
            if l.student_id not in seen:
                seen.add(l.student_id)
                uniq.append(l)
        scores = student_scores([l.student_id for l in uniq], sem)
        bands = score_bands()
        counts = [0] * len(bands)
        rows = []
        for l in uniq:
            s = scores[l.student_id]
            idx, label = band_of(s['score'])
            counts[idx] += 1
            rows.append({'student': l.student, 'room': l.classroom, **s, 'band': label})
        total = len(rows)
        band_rows = [{'label': b[0], 'count': counts[i], 'pct': (counts[i] * 100.0 / total) if total else 0}
                     for i, b in enumerate(bands)]
        return sem, rooms, rows, band_rows, total

    @app.route('/behavior/report')
    @login_required
    @role_required('teacher', 'admin')
    def behavior_report():
        rooms = active_rooms()
        scope = request.args.get('room_id', '')
        if scope != 'all':
            scope = request.args.get('room_id', type=int) or (advisor_rooms()[0].id if advisor_rooms() else 'all')
        sem, sel_rooms, rows, band_rows, total = _report_data(scope)
        # อันดับ (Top N) ตามประเภท/รายการ — ปรับจากตัวกรองของ SSS
        top_kind = request.args.get('top_kind', 'minus')
        top_item = request.args.get('top_item', type=int)
        top_n = request.args.get('top_n', type=int) or 10
        tq = scoped(BehaviorRecord).filter(BehaviorRecord.semester_id == (sem.id if sem else None))
        if scope != 'all':
            tq = tq.filter(BehaviorRecord.classroom_id == scope)
        if top_kind in ('plus', 'minus'):
            tq = tq.filter(BehaviorRecord.kind == top_kind)
        if top_item:
            tq = tq.filter(BehaviorRecord.item_id == top_item)
        agg = {}
        for r in tq.all():
            a = agg.setdefault(r.student_id, {'student': r.student, 'room': r.classroom, 'times': 0, 'points': 0})
            a['times'] += 1
            a['points'] += abs(r.points)
        top = sorted(agg.values(), key=lambda a: (-a['points'], -a['times']))[:top_n]
        items = scoped(BehaviorItem).order_by(BehaviorItem.kind, BehaviorItem.points).all()
        rows.sort(key=lambda r: (r['room'].name if r['room'] else '', r['student'].username or ''))
        return render_template('sss/behavior_report.html', rooms=rooms, scope=scope, sel_rooms=sel_rooms,
                               rows=rows, band_rows=band_rows, total=total, sem=sem, base=base_score(),
                               top=top, top_kind=top_kind, top_item=top_item, top_n=top_n, items=items)

    @app.route('/behavior/report.xlsx')
    @login_required
    @role_required('teacher', 'admin')
    def behavior_report_xlsx():
        from openpyxl import Workbook
        scope = request.args.get('room_id', 'all')
        if scope != 'all':
            scope = request.args.get('room_id', type=int) or 'all'
        sem, sel_rooms, rows, band_rows, total = _report_data(scope)
        wb = Workbook()
        ws = wb.active
        ws.title = 'คะแนนความประพฤติ'
        ws.append([f'รายงานคะแนนความประพฤติ {sem.name if sem else ""}'])
        ws.append(['ที่', 'ห้อง', 'รหัส', 'ชื่อ-สกุล', 'คะแนนตั้งต้น', 'เพิ่ม', 'ตัด', 'คงเหลือ', 'ช่วงคะแนน', 'จำนวนครั้ง'])
        for i, r in enumerate(rows, 1):
            ws.append([i, r['room'].name if r['room'] else '', r['student'].username, r['student'].full_name,
                       base_score(), r['plus'], r['minus'], r['score'], r['band'], r['count']])
        ws2 = wb.create_sheet('สรุปช่วงคะแนน')
        ws2.append(['ช่วงคะแนน', 'จำนวน(คน)', 'ร้อยละ'])
        for b in band_rows:
            ws2.append([b['label'], b['count'], round(b['pct'], 2)])
        ws2.append(['รวม', total, 100 if total else 0])
        name = 'ทุกห้อง' if scope == 'all' else (sel_rooms[0].name.replace('/', '-') if sel_rooms else 'ห้อง')
        return xlsx_response(wb, f'behavior_{name}.xlsx')

    @app.route('/behavior/settings', methods=['GET', 'POST'])
    @login_required
    @role_required('admin')
    def behavior_settings():
        ensure_behavior_items()
        if request.method == 'POST':
            action = request.form.get('action')
            if action == 'config':
                try:
                    cfg_set('behavior_base_score', int(request.form.get('base_score', DEFAULT_BASE_SCORE)))
                    edges = sorted({int(x) for x in (request.form.get('bands') or '').split(',') if x.strip()})
                    if not edges:
                        raise ValueError
                    cfg_set('behavior_bands', ','.join(str(x) for x in edges))
                    db.session.commit()
                    flash('บันทึกการตั้งค่าแล้ว', 'success')
                except ValueError:
                    db.session.rollback()
                    flash('กรุณากรอกตัวเลขให้ถูกต้อง เช่น 0,50,70,90,100', 'danger')
            elif action == 'add':
                name = (request.form.get('name') or '').strip()
                kind = request.form.get('kind')
                pts = request.form.get('points', type=int)
                if name and kind in ('plus', 'minus') and pts and pts > 0:
                    db.session.add(BehaviorItem(name=name[:255], kind=kind, points=pts, is_active=True,
                                                sort_order=scoped(BehaviorItem).count(), school_id=sid()))
                    db.session.commit()
                    flash('เพิ่มรายการแล้ว', 'success')
                else:
                    flash('กรอกชื่อรายการ ประเภท และคะแนน (มากกว่า 0) ให้ครบ', 'danger')
            elif action == 'update':
                item = db.session.get(BehaviorItem, request.form.get('item_id', type=int) or -1)
                if item and same_school(item):
                    name = (request.form.get('name') or '').strip()
                    pts = request.form.get('points', type=int)
                    if name:
                        item.name = name[:255]
                    if pts and pts > 0:
                        item.points = pts
                    item.is_active = request.form.get('is_active') == '1'
                    db.session.commit()
                    flash('บันทึกรายการแล้ว', 'success')
            return redirect(url_for('behavior_settings'))
        items = scoped(BehaviorItem).order_by(BehaviorItem.kind.desc(), BehaviorItem.points, BehaviorItem.sort_order).all()
        return render_template('sss/behavior_settings.html', items=items, base=base_score(),
                               bands=','.join(str(x) for x in band_edges()), band_list=score_bands())

    # ------------------------------------------------------------------
    # [2] เช็กชื่อโฮมรูม
    # ------------------------------------------------------------------
    def ensure_homeroom_activity(room_id):
        act = ClassroomActivity.query.filter_by(classroom_id=room_id, activity_type='homeroom').first()
        if not act:
            # is_active=False เพื่อไม่ให้ไปปนในรายการกิจกรรมพิเศษของห้อง (จัดการผ่านหน้าโฮมรูมเท่านั้น)
            act = ClassroomActivity(classroom_id=room_id, title='กิจกรรมโฮมรูม', activity_type='homeroom',
                                    target_scope='classroom', event_date=None,
                                    note='เช็กชื่อกิจกรรมโฮมรูม (จัดการที่เมนู เช็กชื่อโฮมรูม)', is_active=False)
            db.session.add(act)
            db.session.commit()
        return act

    @app.route('/homeroom', methods=['GET', 'POST'])
    @login_required
    @role_required('teacher', 'admin')
    def homeroom_check():
        rooms = advisor_rooms()
        if not rooms:
            flash('ยังไม่มีห้องที่ปรึกษาของคุณ ให้แอดมินผูกครูประจำชั้นที่หน้า ห้องเรียน ก่อน', 'warning')
            return render_template('sss/homeroom.html', rooms=[], room=None)
        room = pick_room(rooms, request.values.get('room_id', type=int))
        act = ensure_homeroom_activity(room.id)
        sel_date = parse_date(request.values.get('date'), local_today())
        links = room_links(room.id)

        if request.method == 'POST':
            if request.form.get('action') == 'delete':
                ActivityAttendance.query.filter_by(activity_id=act.id, date=sel_date).delete()
                HomeroomLog.query.filter_by(classroom_id=room.id, log_date=sel_date).delete()
                db.session.commit()
                flash(f'ลบข้อมูลโฮมรูม {thai_date_label(sel_date)} แล้ว', 'success')
                return redirect(url_for('homeroom_check', room_id=room.id, date=sel_date.isoformat()))
            log = HomeroomLog.query.filter_by(classroom_id=room.id, log_date=sel_date).first()
            if not log:
                log = HomeroomLog(classroom_id=room.id, log_date=sel_date, school_id=sid())
                db.session.add(log)
            log.topic = (request.form.get('topic') or '').strip()
            log.note = (request.form.get('note') or '').strip()
            log.recorded_by_id = current_user.id
            existing = {a.student_id: a for a in ActivityAttendance.query.filter_by(activity_id=act.id, date=sel_date).all()}
            for l in links:
                st = normalize_attendance_status(request.form.get(f's_{l.student_id}', 'มา'))
                row = existing.get(l.student_id)
                if not row:
                    row = ActivityAttendance(activity_id=act.id, classroom_id=room.id, student_id=l.student_id, date=sel_date)
                    db.session.add(row)
                row.status = st
            db.session.commit()
            flash(f'บันทึกเช็กชื่อโฮมรูม {thai_date_label(sel_date)} แล้ว', 'success')
            return redirect(url_for('homeroom_check', room_id=room.id, date=sel_date.isoformat()))

        old = {a.student_id: normalize_attendance_status(a.status)
               for a in ActivityAttendance.query.filter_by(activity_id=act.id, date=sel_date).all()}
        log = HomeroomLog.query.filter_by(classroom_id=room.id, log_date=sel_date).first()
        sem = get_active_semester()
        start, end = sem_range(sem)
        logs = (HomeroomLog.query.filter_by(classroom_id=room.id)
                .filter(HomeroomLog.log_date >= start, HomeroomLog.log_date <= end)
                .order_by(HomeroomLog.log_date.desc()).limit(30).all())
        counts = {}
        for a in ActivityAttendance.query.filter(ActivityAttendance.activity_id == act.id,
                                                 ActivityAttendance.date.in_([x.log_date for x in logs] or [date(1900, 1, 1)])).all():
            c = counts.setdefault(a.date, {})
            st = normalize_attendance_status(a.status)
            c[st] = c.get(st, 0) + 1
        return render_template('sss/homeroom.html', rooms=rooms, room=room, links=links, old=old, log=log,
                               sel_date=sel_date, statuses=HR_STATUSES, logs=logs, counts=counts,
                               is_blocked=is_school_blocked_day(sel_date, room.teacher_id),
                               already_saved=bool(old))

    # ------------------------------------------------------------------
    # [4] รายงานเช็กชื่อกิจกรรม (ตารางนักเรียน × วัน) + ภาพรวมรายวัน
    # ------------------------------------------------------------------
    def build_grid(room, rtype, start, end):
        links = room_links(room.id)
        q = (ActivityAttendance.query
             .join(ClassroomActivity, ActivityAttendance.activity_id == ClassroomActivity.id)
             .filter(ActivityAttendance.classroom_id == room.id,
                     ActivityAttendance.date >= start, ActivityAttendance.date <= end))
        if rtype == 'special':
            q = q.filter(ClassroomActivity.activity_type.notin_(['assembly', 'homeroom']))
        else:
            q = q.filter(ClassroomActivity.activity_type == rtype)
        rows = q.order_by(ActivityAttendance.date.asc()).all()
        topics = {}
        if rtype == 'homeroom':
            topics = {h.log_date: h.topic for h in HomeroomLog.query.filter(
                HomeroomLog.classroom_id == room.id, HomeroomLog.log_date >= start, HomeroomLog.log_date <= end).all()}
        cols = {}
        data = {}
        for a in rows:
            if rtype == 'special':
                key = f'{a.date.isoformat()}#{a.activity_id}'
                title = a.activity.title if a.activity else ''
            else:
                key = a.date.isoformat()
                title = topics.get(a.date, '')
            cols.setdefault(key, {'key': key, 'date': a.date, 'short': a.date.strftime('%d/%m/') + str(a.date.year + 543)[-2:],
                                  'title': title})
            data.setdefault(a.student_id, {})[key] = normalize_attendance_status(a.status)
        columns = sorted(cols.values(), key=lambda c: (c['date'], c['key']))
        summary = []
        for l in links:
            m = data.get(l.student_id, {})
            c = {'มา': 0, 'สาย': 0, 'ขาด': 0, 'ลาป่วย': 0, 'โดดเรียน': 0, 'ไปกิจกรรม': 0}
            for st in m.values():
                c[st] = c.get(st, 0) + 1
            checked = len(m)
            attended = sum(c.get(s, 0) for s in ATTENDED_STATUSES)
            summary.append({'student': l.student, 'cells': m, 'c': c, 'checked': checked,
                            'pct': (attended * 100.0 / checked) if checked else None})
        return columns, summary

    def report_params():
        rooms = my_rooms() if current_user.role != 'admin' else active_rooms()
        room = pick_room(rooms, request.args.get('room_id', type=int))
        rtype = request.args.get('type', 'assembly')
        if rtype not in REPORT_TYPES:
            rtype = 'assembly'
        period = request.args.get('period', 'month')
        sem = get_active_semester()
        s_start, s_end = sem_range(sem)
        month_text = request.args.get('month') or local_today().strftime('%Y-%m')
        try:
            y, m = [int(x) for x in month_text.split('-')]
            m_start = date(y, m, 1)
        except ValueError:
            m_start = local_today().replace(day=1)
        m_end = (m_start.replace(day=28) + timedelta(days=4)).replace(day=1) - timedelta(days=1)
        if period == 'term':
            start, end = s_start, s_end
            period_label = f'ทั้งภาคเรียน ({sem.name})' if sem else 'ทั้งปี'
        else:
            period = 'month'
            start, end = m_start, m_end
            period_label = f'เดือน{THAI_MONTHS[m_start.month]} {m_start.year + 543}'
        months = []
        cur = s_start.replace(day=1)
        while cur <= s_end and len(months) < 24:
            months.append({'value': cur.strftime('%Y-%m'), 'label': f'{THAI_MONTHS[cur.month]} {cur.year + 543}'})
            cur = (cur.replace(day=28) + timedelta(days=4)).replace(day=1)
        sel_val = m_start.strftime('%Y-%m')
        if sel_val not in [m['value'] for m in months]:
            months.insert(0, {'value': sel_val, 'label': f'{THAI_MONTHS[m_start.month]} {m_start.year + 543}'})
        return rooms, room, rtype, period, start, end, period_label, months, m_start.strftime('%Y-%m'), sem

    @app.route('/activity-report')
    @login_required
    @role_required('teacher', 'admin')
    def activity_report():
        rooms, room, rtype, period, start, end, period_label, months, month_val, sem = report_params()
        columns, summary, hr_logs, averages = [], [], [], None
        advisors = []
        if room:
            columns, summary = build_grid(room, rtype, start, end)
            advisors = [l.teacher for l in TeacherClassroom.query.filter_by(classroom_id=room.id).all() if l.teacher]
            if room.teacher and room.teacher not in advisors:
                advisors.insert(0, room.teacher)
            if rtype == 'homeroom':
                act = ClassroomActivity.query.filter_by(classroom_id=room.id, activity_type='homeroom').first()
                logs = (HomeroomLog.query.filter(HomeroomLog.classroom_id == room.id,
                                                 HomeroomLog.log_date >= start, HomeroomLog.log_date <= end)
                        .order_by(HomeroomLog.log_date).all())
                names = {}
                if act:
                    for a in ActivityAttendance.query.filter(ActivityAttendance.activity_id == act.id,
                                                             ActivityAttendance.date >= start,
                                                             ActivityAttendance.date <= end).all():
                        st = normalize_attendance_status(a.status)
                        if st not in ATTENDED_STATUSES:
                            names.setdefault(a.date, []).append(f'{a.student.full_name if a.student else "?"} ({status_label(st)})')
                hr_logs = [{'log': lg, 'absent': names.get(lg.log_date, [])} for lg in logs]
            total_cells = sum(r['checked'] for r in summary)
            if total_cells:
                averages = {
                    'มา': sum(r['c'].get('มา', 0) + r['c'].get('สาย', 0) + r['c'].get('ไปกิจกรรม', 0) for r in summary) * 100.0 / total_cells,
                    'ลา': sum(r['c'].get('ลาป่วย', 0) for r in summary) * 100.0 / total_cells,
                    'ขาด': sum(r['c'].get('ขาด', 0) + r['c'].get('โดดเรียน', 0) for r in summary) * 100.0 / total_cells,
                }
        return render_template('sss/activity_report.html', rooms=rooms, room=room, rtype=rtype, types=REPORT_TYPES,
                               period=period, period_label=period_label, months=months, month_val=month_val,
                               columns=columns, summary=summary, symbols=ATTENDANCE_SYMBOLS, sem=sem,
                               hr_logs=hr_logs, averages=averages, advisors=advisors)

    @app.route('/activity-report.xlsx')
    @login_required
    @role_required('teacher', 'admin')
    def activity_report_xlsx():
        from openpyxl import Workbook
        from openpyxl.styles import Alignment
        rooms, room, rtype, period, start, end, period_label, months, month_val, sem = report_params()
        if not room:
            abort(404)
        columns, summary = build_grid(room, rtype, start, end)
        wb = Workbook()
        ws = wb.active
        ws.title = 'รายงาน'
        ws.append([f'รายงานเช็กชื่อ{REPORT_TYPES[rtype]} ห้อง {room.name} {period_label}'])
        head = ['ที่', 'รหัส', 'ชื่อ-สกุล'] + [c['short'] + (f"\n{c['title']}" if c['title'] else '') for c in columns] \
            + ['มา', 'สาย', 'ขาด', 'ลา', 'อื่น ๆ', 'เช็กทั้งหมด', '% เข้าร่วม']
        ws.append(head)
        for cell in ws[2]:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        for i, r in enumerate(summary, 1):
            c = r['c']
            other = c.get('โดดเรียน', 0) + c.get('ไปกิจกรรม', 0)
            ws.append([i, r['student'].username, r['student'].full_name]
                      + [ATTENDANCE_SYMBOLS.get(r['cells'].get(col['key']), '') for col in columns]
                      + [c.get('มา', 0), c.get('สาย', 0), c.get('ขาด', 0), c.get('ลาป่วย', 0), other, r['checked'],
                         round(r['pct'], 2) if r['pct'] is not None else ''])
        ws.append([])
        ws.append(['สัญลักษณ์: ' + '  '.join(f'{v}={status_label(k)}' for k, v in ATTENDANCE_SYMBOLS.items())])
        return xlsx_response(wb, f'activity_{rtype}_{room.name.replace("/", "-")}.xlsx')

    @app.route('/activity-report/daily')
    @login_required
    @role_required('teacher', 'admin')
    def activity_report_daily():
        sel_date = parse_date(request.args.get('date'), local_today())
        rooms = active_rooms()
        acts = ClassroomActivity.query.filter(ClassroomActivity.activity_type.in_(['assembly', 'homeroom']),
                                              ClassroomActivity.classroom_id.in_([r.id for r in rooms] or [-1])).all()
        act_type = {a.id: a.activity_type for a in acts}
        rows_db = ActivityAttendance.query.filter(ActivityAttendance.date == sel_date,
                                                  ActivityAttendance.activity_id.in_(list(act_type) or [-1])).all()
        stat = {}
        for a in rows_db:
            t = act_type.get(a.activity_id)
            st = normalize_attendance_status(a.status)
            d = stat.setdefault((a.classroom_id, t), {})
            d[st] = d.get(st, 0) + 1
        out, totals = [], {'assembly': {}, 'homeroom': {}}
        for r in rooms:
            n = ClassroomStudent.query.filter_by(classroom_id=r.id).count()
            row = {'room': r, 'n': n}
            for t in ('assembly', 'homeroom'):
                d = stat.get((r.id, t), {})
                row[t] = d
                for k, v in d.items():
                    totals[t][k] = totals[t].get(k, 0) + v
            out.append(row)
        return render_template('sss/activity_daily.html', sel_date=sel_date, rows=out, totals=totals,
                               blocked=is_school_blocked_day(sel_date))

    # ------------------------------------------------------------------
    # [7] บันทึกอนุญาตออกนอกบริเวณโรงเรียน
    # ------------------------------------------------------------------
    def leave_stats():
        today = local_today()
        start_today = datetime.combine(today, datetime.min.time())
        Q = lambda: scoped(CampusLeave)
        start_month = datetime.combine(today.replace(day=1), datetime.min.time())
        sem = get_active_semester()
        s_start, _ = sem_range(sem)
        return {
            'today': Q().filter(CampusLeave.out_at >= start_today).count(),
            'month': Q().filter(CampusLeave.out_at >= start_month).count(),
            'term': Q().filter(CampusLeave.out_at >= datetime.combine(s_start, datetime.min.time())).count(),
            'out_now': Q().filter(CampusLeave.returned_at.is_(None)).count(),
        }

    @app.route('/campus-leave', methods=['GET', 'POST'])
    @login_required
    @role_required('teacher', 'admin')
    def campus_leave():
        if request.method == 'POST':
            student = school_student(request.form.get('student_id', type=int))
            dest = (request.form.get('destination') or '').strip()
            if not student or not dest:
                flash('กรุณาเลือกนักเรียนและกรอกสถานที่ที่ขอไป', 'danger')
                return redirect(url_for('campus_leave', q=request.form.get('q') or None))
            still_out = scoped(CampusLeave).filter_by(student_id=student.id, returned_at=None).first()
            if still_out:
                flash(f'{student.full_name} ยังไม่ได้บันทึกกลับเข้าโรงเรียนจากรายการก่อนหน้า '
                      f'(ออกเมื่อ {still_out.out_at.strftime("%H:%M")}) กรุณากด "บันทึกกลับเข้า" ก่อน', 'warning')
                return redirect(url_for('campus_leave'))
            now = local_now().replace(tzinfo=None, second=0, microsecond=0)
            out_time = (request.form.get('out_time') or '').strip()
            out_at = now
            if out_time:
                try:
                    hh, mm = [int(x) for x in out_time.split(':')[:2]]
                    out_at = now.replace(hour=hh, minute=mm)
                except ValueError:
                    pass
            room = student_room(student.id)
            sem = get_active_semester()
            db.session.add(CampusLeave(
                student_id=student.id, classroom_id=room.id if room else None, destination=dest[:255],
                reason=(request.form.get('reason') or '').strip(), out_at=out_at,
                expect_return=(request.form.get('expect_return') or '').strip()[:20],
                note=(request.form.get('note') or '').strip(), approved_by_id=current_user.id,
                semester_id=sem.id if sem else None, school_id=sid()))
            db.session.commit()
            flash(f'บันทึกอนุญาตให้ {student.full_name} ออกนอกบริเวณแล้ว', 'success')
            return redirect(url_for('campus_leave'))
        q = (request.args.get('q') or '').strip()
        found = search_students(q, 20) if q else []
        rooms_of = {}
        if found:
            for link in ClassroomStudent.query.filter(ClassroomStudent.student_id.in_([s.id for s in found])).all():
                rooms_of.setdefault(link.student_id, link.classroom.name if link.classroom else '')
        chosen = None
        if request.args.get('student_id', type=int):
            chosen = school_student(request.args.get('student_id', type=int))
        out_now = scoped(CampusLeave).filter(CampusLeave.returned_at.is_(None)).order_by(CampusLeave.out_at.desc()).all()
        today_start = datetime.combine(local_today(), datetime.min.time())
        today_rows = scoped(CampusLeave).filter(CampusLeave.out_at >= today_start).order_by(CampusLeave.out_at.desc()).all()
        return render_template('sss/campus_leave.html', q=q, found=found, rooms_of=rooms_of, chosen=chosen,
                               chosen_room=student_room(chosen.id) if chosen else None,
                               out_now=out_now, today_rows=today_rows, stats=leave_stats(),
                               now_time=local_now().strftime('%H:%M'), today=local_today())

    @app.route('/campus-leave/<int:leave_id>/return', methods=['POST'])
    @login_required
    @role_required('teacher', 'admin')
    def campus_leave_return(leave_id):
        row = scoped(CampusLeave).filter(CampusLeave.id == leave_id).first_or_404()
        if row.returned_at is None:
            row.returned_at = local_now().replace(tzinfo=None, second=0, microsecond=0)
            row.returned_by_id = current_user.id
            db.session.commit()
            flash(f'บันทึกเวลากลับเข้าโรงเรียนของ {row.student.full_name if row.student else ""} แล้ว', 'success')
        return redirect(request.referrer or url_for('campus_leave'))

    @app.route('/campus-leave/<int:leave_id>/delete', methods=['POST'])
    @login_required
    @role_required('teacher', 'admin')
    def campus_leave_delete(leave_id):
        row = scoped(CampusLeave).filter(CampusLeave.id == leave_id).first_or_404()
        if current_user.role != 'admin' and row.approved_by_id != current_user.id:
            flash('ลบได้เฉพาะรายการที่คุณบันทึกเอง', 'danger')
        else:
            db.session.delete(row)
            db.session.commit()
            flash('ลบรายการแล้ว', 'success')
        return redirect(request.referrer or url_for('campus_leave'))

    @app.route('/campus-leave/history')
    @login_required
    @role_required('teacher', 'admin')
    def campus_leave_history():
        sem = get_active_semester()
        s_start, s_end = sem_range(sem)
        d_from = parse_date(request.args.get('date_from'), s_start)
        d_to = parse_date(request.args.get('date_to'), local_today())
        room_id = request.args.get('room_id', type=int)
        q = scoped(CampusLeave).filter(CampusLeave.out_at >= datetime.combine(d_from, datetime.min.time()),
                                     CampusLeave.out_at < datetime.combine(d_to + timedelta(days=1), datetime.min.time()))
        if room_id:
            q = q.filter(CampusLeave.classroom_id == room_id)
        rows = q.order_by(CampusLeave.out_at.desc()).limit(1000).all()
        if request.args.get('export') == 'xlsx':
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.title = 'ออกนอกบริเวณ'
            ws.append(['วันที่', 'เวลาออก', 'เวลากลับ', 'รหัส', 'ชื่อ-สกุล', 'ห้อง', 'ไปที่', 'เหตุผล', 'ผู้อนุญาต', 'หมายเหตุ'])
            for r in rows:
                ws.append([r.out_at.strftime('%d/%m/') + str(r.out_at.year + 543), r.out_at.strftime('%H:%M'),
                           r.returned_at.strftime('%H:%M') if r.returned_at else 'ยังไม่กลับ',
                           r.student.username if r.student else '', r.student.full_name if r.student else '',
                           r.classroom.name if r.classroom else '', r.destination, r.reason,
                           r.approved_by.full_name if r.approved_by else '', r.note])
            return xlsx_response(wb, 'campus_leave.xlsx')
        return render_template('sss/campus_leave_history.html', rows=rows, d_from=d_from, d_to=d_to,
                               rooms=active_rooms(), room_id=room_id)

    return g
