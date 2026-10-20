"""ระบบจัดตารางสอน (Timetable Builder)

- แผนตารางสอน (plan) แยกตามโรงเรียน ทำได้หลายแผน/ลองหลายแบบ
- ข้อมูล: ครู (คาบสูงสุด/วัน), รายวิชา-ห้อง (คาบ/สัปดาห์, คาบติดกัน), ข้อห้าม (ครู/ห้อง/ทั้งโรงเรียน; ห้ามเด็ดขาด/หลีกเลี่ยง)
- จัดอัตโนมัติ 2 โหมด: จัดใหม่ทั้งหมด / แก้เฉพาะจุดที่มีปัญหา (คงตารางเดิมไว้ให้มากที่สุด) — ช่องที่ล็อกไว้จะไม่ถูกย้าย
- ปรับแก้เองบนตาราง (คลิกเลือก-คลิกวาง), ล็อก/ปลดล็อก, ตรวจการชนอัตโนมัติ
- ภาระงานสอน: สรุปคาบต่อครู/ต่อวัน + ข้อเสนอโยกรายวิชาเพื่อเกลี่ยภาระงานในกลุ่มสาระเดียวกัน
- นำเข้า/ส่งออก Excel และพิมพ์ตารางรายครู/รายห้อง

เรียกใช้: ``timetable_builder.register(app, globals())`` ใน app.py ก่อน init_db()
"""
import re
import json
import difflib
from io import BytesIO
from datetime import datetime
from collections import defaultdict

from flask import render_template, request, redirect, url_for, flash, abort, send_file, jsonify
from flask_login import login_required, current_user

import timetable_solver as TS

DAY_NAMES = ['จันทร์', 'อังคาร', 'พุธ', 'พฤหัสบดี', 'ศุกร์', 'เสาร์', 'อาทิตย์']
DEFAULT_TIMES = '08:40-09:30,09:31-10:20,10:31-11:20,11:21-12:10,13:01-13:50,13:51-14:40,14:41-15:30,15:31-16:20'
SUBJECT_GROUPS = {
    'ท': 'ภาษาไทย', 'ค': 'คณิตศาสตร์', 'ว': 'วิทยาศาสตร์และเทคโนโลยี', 'ส': 'สังคมศึกษาฯ',
    'พ': 'สุขศึกษาและพลศึกษา', 'ศ': 'ศิลปะ', 'ง': 'การงานอาชีพ', 'อ': 'ภาษาต่างประเทศ', 'จ': 'ภาษาต่างประเทศ',
    'ญ': 'ภาษาต่างประเทศ', 'ก': 'กิจกรรมพัฒนาผู้เรียน', 'น': 'กิจกรรมพัฒนาผู้เรียน',
}
SCHEDULE_HEADERS = ['ครู', 'วัน', 'คาบ', 'รหัสวิชา', 'ชื่อวิชา', 'ชั้น/ห้อง', 'ห้องเรียน', 'หัวข้อ']
PALETTE = ['#dbeafe', '#dcfce7', '#fef3c7', '#fce7f3', '#ede9fe', '#cffafe', '#fee2e2', '#e0e7ff',
           '#fae8ff', '#ecfccb', '#ffedd5', '#e2e8f0']


def subject_group_of(code):
    code = (code or '').strip()
    return SUBJECT_GROUPS.get(code[:1], 'อื่น ๆ') if code else 'อื่น ๆ'


def expand_class(label):
    """'ม.1/2-3' -> ['ม.1/2','ม.1/3'], 'ม.4/1-3' -> 3 ห้อง, 'ม.1/1' -> ['ม.1/1'], 'ม.1/1,ม.2/1' -> 2 ห้อง"""
    out = []
    for part in re.split(r'[,،;]+', label or ''):
        part = part.strip().replace(' ', '')
        if not part:
            continue
        m = re.match(r'^(.*?)(\d+)/(\d+)-(\d+)$', part)
        if m:
            pre, g, a, b = m.group(1), m.group(2), int(m.group(3)), int(m.group(4))
            if a <= b <= a + 20:
                out.extend(f'{pre}{g}/{i}' for i in range(a, b + 1))
                continue
        out.append(part)
    seen, res = set(), []
    for x in out:
        if x not in seen:
            seen.add(x)
            res.append(x)
    return res


def parse_day(v):
    v = str(v or '').strip()
    if v.isdigit():
        return int(v)
    for i, n in enumerate(DAY_NAMES, 1):
        if v.startswith(n) or v == n[:2] or v.replace('วัน', '').startswith(n):
            return i
    return None


def parse_periods(v, maxp):
    """'1-4' / '1,2' / 'ทั้งวัน' / 3 -> list[int]"""
    v = str(v or '').strip()
    if not v or 'ทั้งวัน' in v or v in ('*', 'all'):
        return list(range(1, maxp + 1))
    if 'เช้า' in v:
        return None  # ให้ผู้เรียกแทนด้วยช่วงเช้า
    if 'บ่าย' in v:
        return None
    out = []
    for part in re.split(r'[,\s]+', v):
        if not part:
            continue
        m = re.match(r'^(\d+)\s*-\s*(\d+)$', part)
        if m:
            out.extend(range(int(m.group(1)), int(m.group(2)) + 1))
        elif part.isdigit():
            out.append(int(part))
    return [p for p in out if 1 <= p <= maxp]


def register(app, g):
    db = g['db']
    User = g['User']
    Subject = g['Subject']
    Classroom = g['Classroom']
    TeacherSubject = g['TeacherSubject']
    TeacherClassroom = g['TeacherClassroom']
    SubjectClassroom = g['SubjectClassroom']
    TeachingSchedule = g['TeachingSchedule']
    PeriodLessonLog = g['PeriodLessonLog']
    Attendance = g['Attendance']
    get_active_semester = g['get_active_semester']
    role_required = g['role_required']
    _current_school_id = g.get('current_school_id')
    local_now = g.get('local_now') or datetime.now

    def sid():
        try:
            return _current_school_id() if _current_school_id else getattr(current_user, 'school_id', None)
        except Exception:
            return getattr(current_user, 'school_id', None)

    # ------------------------------------------------------------------
    # Models
    # ------------------------------------------------------------------
    class TtPlan(db.Model):
        __tablename__ = 'tt_plan'
        id = db.Column(db.Integer, primary_key=True)
        school_id = db.Column(db.Integer, nullable=True, index=True)
        name = db.Column(db.String(200), nullable=False)
        days = db.Column(db.Integer, default=5)
        periods = db.Column(db.Integer, default=8)
        lunch_after = db.Column(db.Integer, default=4)
        period_times = db.Column(db.Text, default=DEFAULT_TIMES)
        default_max_per_day = db.Column(db.Integer, default=6)
        max_consecutive = db.Column(db.Integer, default=4)
        note = db.Column(db.Text, default='')
        last_cost = db.Column(db.Integer, nullable=True)
        last_run_at = db.Column(db.DateTime, nullable=True)
        created_by_id = db.Column(db.Integer, nullable=True)
        created_at = db.Column(db.DateTime, default=datetime.utcnow)

    class TtTeacher(db.Model):
        __tablename__ = 'tt_teacher'
        id = db.Column(db.Integer, primary_key=True)
        plan_id = db.Column(db.Integer, db.ForeignKey('tt_plan.id'), nullable=False, index=True)
        name = db.Column(db.String(200), nullable=False)
        user_id = db.Column(db.Integer, nullable=True)
        max_per_day = db.Column(db.Integer, nullable=True)
        subject_groups = db.Column(db.Text, default='')   # กลุ่มสาระที่สอนได้ (คั่นด้วย ,)
        room = db.Column(db.String(120), default='')

    class TtLesson(db.Model):
        __tablename__ = 'tt_lesson'
        id = db.Column(db.Integer, primary_key=True)
        plan_id = db.Column(db.Integer, db.ForeignKey('tt_plan.id'), nullable=False, index=True)
        teacher_id = db.Column(db.Integer, db.ForeignKey('tt_teacher.id'), nullable=True, index=True)
        code = db.Column(db.String(40), default='')
        name = db.Column(db.String(200), default='')
        class_label = db.Column(db.String(120), default='')
        groups = db.Column(db.Text, default='')
        periods = db.Column(db.Integer, default=1)
        block = db.Column(db.Integer, default=1)
        room = db.Column(db.String(120), default='')
        teacher = db.relationship('TtTeacher')

        @property
        def group_list(self):
            return [x for x in (self.groups or '').split(',') if x]

        @property
        def label(self):
            return f'{self.code} {self.name}'.strip()

    class TtBlock(db.Model):
        __tablename__ = 'tt_block'
        id = db.Column(db.Integer, primary_key=True)
        plan_id = db.Column(db.Integer, db.ForeignKey('tt_plan.id'), nullable=False, index=True)
        scope = db.Column(db.String(10), default='teacher')   # teacher / group / all
        teacher_id = db.Column(db.Integer, nullable=True)
        group_name = db.Column(db.String(60), default='')
        day = db.Column(db.Integer, nullable=False)
        period = db.Column(db.Integer, nullable=False)
        hard = db.Column(db.Boolean, default=True)
        weight = db.Column(db.Integer, default=30)
        reason = db.Column(db.String(200), default='')

    class TtSlot(db.Model):
        __tablename__ = 'tt_slot'
        id = db.Column(db.Integer, primary_key=True)
        plan_id = db.Column(db.Integer, db.ForeignKey('tt_plan.id'), nullable=False, index=True)
        lesson_id = db.Column(db.Integer, db.ForeignKey('tt_lesson.id'), nullable=False, index=True)
        day = db.Column(db.Integer, nullable=False)
        period = db.Column(db.Integer, nullable=False)
        locked = db.Column(db.Boolean, default=False)
        lesson = db.relationship('TtLesson')

    for _c in (TtPlan, TtTeacher, TtLesson, TtBlock, TtSlot):
        g[_c.__name__] = _c

    # ------------------------------------------------------------------
    # helpers
    # ------------------------------------------------------------------
    def get_plan(pid):
        p = db.session.get(TtPlan, pid)
        if not p or p.school_id != sid():
            abort(404)
        return p

    def times_of(plan):
        t = [x.strip() for x in (plan.period_times or '').split(',')]
        return [(t[i] if i < len(t) else '') for i in range(plan.periods)]

    def plan_data(plan):
        teachers = TtTeacher.query.filter_by(plan_id=plan.id).order_by(TtTeacher.name).all()
        lessons = TtLesson.query.filter_by(plan_id=plan.id).order_by(TtLesson.code, TtLesson.class_label).all()
        blocks = TtBlock.query.filter_by(plan_id=plan.id).all()
        slots = TtSlot.query.filter_by(plan_id=plan.id).all()
        return teachers, lessons, blocks, slots

    def build_problem(plan, teachers, lessons, blocks, slots, locked_only=True):
        tmap = {t.id: {'max_per_day': t.max_per_day or plan.default_max_per_day or 0} for t in teachers}
        L = [{'id': l.id, 'teacher_id': l.teacher_id or -l.id, 'groups': l.group_list or [f'#{l.id}'],
              'periods': max(0, int(l.periods or 0)), 'block': max(1, int(l.block or 1))}
             for l in lessons if (l.periods or 0) > 0]
        B = []
        for b in blocks:
            ref = b.teacher_id if b.scope == 'teacher' else (b.group_name if b.scope == 'group' else None)
            B.append({'scope': b.scope, 'ref': ref, 'day': b.day, 'period': b.period, 'hard': bool(b.hard),
                      'weight': b.weight or 30, 'reason': b.reason or ''})
        valid = {l['id'] for l in L}
        locked = [{'lesson_id': s.lesson_id, 'day': s.day, 'period': s.period}
                  for s in slots if s.locked and s.lesson_id in valid]
        return {'days': plan.days, 'periods': plan.periods, 'lunch_after': plan.lunch_after,
                'max_consecutive': plan.max_consecutive or 4, 'teachers': tmap, 'lessons': L,
                'blocks': B, 'locked': locked}

    def analyze(plan, teachers, lessons, blocks, slots):
        prob = build_problem(plan, teachers, lessons, blocks, slots)
        placements = [{'lesson_id': s.lesson_id, 'day': s.day, 'period': s.period} for s in slots]
        issues = TS.evaluate(prob, placements)
        placed = defaultdict(int)
        for s in slots:
            placed[s.lesson_id] += 1
        missing = [(l, l.periods - placed.get(l.id, 0)) for l in lessons if placed.get(l.id, 0) < (l.periods or 0)]
        bad_cells = set()
        for i in issues:
            if i['type'] in ('teacher_clash', 'group_clash', 'blocked_hard'):
                for lid in i.get('lessons', [i.get('lesson_id')]):
                    bad_cells.add((lid, i['day'], i['period']))
        soft_cells = {(i['lesson_id'], i['day'], i['period']) for i in issues if i['type'] == 'blocked_soft'}
        return issues, missing, bad_cells, soft_cells

    def teacher_loads(plan, teachers, lessons, slots):
        lmap = {l.id: l for l in lessons}
        per_day = {t.id: [0] * plan.days for t in teachers}
        assigned = defaultdict(int)
        for l in lessons:
            if l.teacher_id:
                assigned[l.teacher_id] += l.periods or 0
        for s in slots:
            l = lmap.get(s.lesson_id)
            if l and l.teacher_id in per_day and 1 <= s.day <= plan.days:
                per_day[l.teacher_id][s.day - 1] += 1
        return per_day, assigned

    def colors_for(lessons):
        codes = sorted({l.code for l in lessons})
        cmap = {c: PALETTE[i % len(PALETTE)] for i, c in enumerate(codes)}
        return {l.id: cmap.get(l.code, '#e2e8f0') for l in lessons}

    def all_groups(lessons):
        def keyf(x):
            m = re.match(r'^\D*?(\d+)/(\d+)$', x)
            return (0, int(m.group(1)), int(m.group(2)), x) if m else (1, 0, 0, x)
        return sorted({g_ for l in lessons for g_ in l.group_list}, key=keyf)

    def xlsx_response(wb, filename):
        bio = BytesIO()
        wb.save(bio)
        bio.seek(0)
        return send_file(bio, as_attachment=True, download_name=filename,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

    def lesson_from_form(plan, lesson, form):
        lesson.code = (form.get('code') or '').strip()[:40]
        lesson.name = (form.get('name') or '').strip()[:200]
        lesson.class_label = (form.get('class_label') or '').strip()[:120]
        lesson.groups = ','.join(expand_class(lesson.class_label))
        lesson.periods = max(0, min(40, int(form.get('periods') or 1)))
        lesson.block = max(1, min(4, int(form.get('block') or 1)))
        lesson.room = (form.get('room') or '').strip()[:120]
        tid = form.get('teacher_id', type=int)
        t = db.session.get(TtTeacher, tid) if tid else None
        lesson.teacher_id = t.id if t and t.plan_id == plan.id else None

    def trim_extra_slots(lesson):
        slots = TtSlot.query.filter_by(lesson_id=lesson.id).order_by(TtSlot.locked.asc(), TtSlot.id.desc()).all()
        extra = len(slots) - (lesson.periods or 0)
        for s in slots[:max(0, extra)]:
            db.session.delete(s)

    def person_key(name):
        """Normalize Thai teacher names while keeping matching conservative."""
        value = (name or '').strip()
        for prefix in ('นางสาว', 'นาย', 'นาง', 'ครู'):
            if value.startswith(prefix):
                value = value[len(prefix):]
                break
        return re.sub(r'[^ก-๙A-Za-z0-9]', '', value).lower()

    def suggested_user(tt_teacher, users):
        if tt_teacher.user_id:
            linked = next((u for u in users if u.id == tt_teacher.user_id), None)
            if linked:
                return linked, 1.0
        key = person_key(tt_teacher.name)
        exact = next((u for u in users if person_key(u.full_name) == key), None)
        if exact:
            return exact, 1.0
        scored = [(difflib.SequenceMatcher(None, key, person_key(u.full_name)).ratio(), u) for u in users]
        score, user = max(scored, default=(0, None), key=lambda x: x[0])
        return (user, score) if score >= .80 else (None, score)

    def teacher_for_user(plan, user):
        """Find the timetable teacher belonging to a login account.

        Older imported plans may not have user_id populated, so retain the
        conservative name matching already used by the timetable module.
        """
        linked = TtTeacher.query.filter_by(plan_id=plan.id, user_id=user.id).first()
        if linked:
            return linked
        key = person_key(user.full_name)
        return next((t for t in TtTeacher.query.filter_by(plan_id=plan.id).all()
                     if person_key(t.name) == key), None)

    @app.context_processor
    def _tt_ctx():
        return {'tt_day_names': DAY_NAMES, 'timetable_enabled': True}

    # ------------------------------------------------------------------
    # หน้าแรก: รายการแผน / สร้าง / นำเข้า
    # ------------------------------------------------------------------
    @app.route('/timetable')
    @login_required
    @role_required('admin')
    def tt_home():
        plans = TtPlan.query.filter_by(school_id=sid()).order_by(TtPlan.created_at.desc()).all()
        stats = {}
        for p in plans:
            stats[p.id] = {
                'teachers': TtTeacher.query.filter_by(plan_id=p.id).count(),
                'lessons': TtLesson.query.filter_by(plan_id=p.id).count(),
                'periods': db.session.query(db.func.coalesce(db.func.sum(TtLesson.periods), 0)).filter(TtLesson.plan_id == p.id).scalar(),
                'placed': TtSlot.query.filter_by(plan_id=p.id).count(),
            }
        return render_template('timetable/home.html', plans=plans, stats=stats, default_times=DEFAULT_TIMES)

    @app.route('/timetable/preferences', methods=['GET', 'POST'])
    @login_required
    @role_required('teacher')
    def tt_preferences():
        """Teacher questionnaire used as constraints for automatic scheduling."""
        plans = TtPlan.query.filter_by(school_id=sid()).order_by(TtPlan.created_at.desc()).all()
        available = []
        for item in plans:
            tt_teacher = teacher_for_user(item, current_user)
            if tt_teacher:
                available.append((item, tt_teacher))

        requested_pid = request.values.get('plan_id', type=int)
        selected = next(((p, t) for p, t in available if not requested_pid or p.id == requested_pid), None)
        if not selected:
            return render_template('timetable/preferences.html', plans=available, plan=None, tt_teacher=None,
                                   lessons=[], current_blocks={}, times=[])

        plan, tt_teacher = selected
        lessons = TtLesson.query.filter_by(plan_id=plan.id, teacher_id=tt_teacher.id).order_by(
            TtLesson.code, TtLesson.class_label).all()

        if request.method == 'POST':
            allowed_lesson_ids = {lesson.id for lesson in lessons}
            for lesson in lessons:
                requested = request.form.get(f'block_{lesson.id}', type=int)
                if lesson.id in allowed_lesson_ids and requested is not None:
                    lesson.block = max(1, min(4, int(lesson.periods or 1), requested))

            # ครูแก้เฉพาะข้อห้ามเด็ดขาดของตัวเอง ข้อเสนอแบบ "หลีกเลี่ยง"
            # ที่ผู้จัดตารางตั้งไว้ยังคงอยู่และไม่ถูกแบบสอบถามนี้ลบ
            TtBlock.query.filter_by(plan_id=plan.id, scope='teacher', teacher_id=tt_teacher.id,
                                    hard=True).delete(synchronize_session=False)
            reason = (request.form.get('reason') or 'ประชุม/ติดภารกิจ').strip()[:200]
            for cell in request.form.getlist('cell'):
                try:
                    day, period = [int(x) for x in cell.split('-', 1)]
                except (TypeError, ValueError):
                    continue
                if 1 <= day <= plan.days and 1 <= period <= plan.periods:
                    db.session.add(TtBlock(plan_id=plan.id, scope='teacher', teacher_id=tt_teacher.id,
                                           day=day, period=period, hard=True, weight=30, reason=reason))
            if not tt_teacher.user_id:
                tt_teacher.user_id = current_user.id
            db.session.commit()
            flash('บันทึกความต้องการจัดตารางสอนแล้ว ระบบจะไม่นำคาบสอนไปลงในช่วงที่เลือก', 'success')
            return redirect(url_for('tt_preferences', plan_id=plan.id))

        current_blocks = {(b.day, b.period): b for b in TtBlock.query.filter_by(
            plan_id=plan.id, scope='teacher', teacher_id=tt_teacher.id, hard=True).all()}
        return render_template('timetable/preferences.html', plans=available, plan=plan, tt_teacher=tt_teacher,
                               lessons=lessons, current_blocks=current_blocks, times=times_of(plan))

    @app.route('/timetable/new', methods=['POST'])
    @login_required
    @role_required('admin')
    def tt_new():
        f = request.form
        name = (f.get('name') or '').strip() or 'ตารางสอนใหม่'
        plan = TtPlan(school_id=sid(), name=name[:200], days=max(1, min(7, f.get('days', 5, type=int))),
                      periods=max(1, min(14, f.get('periods', 8, type=int))),
                      lunch_after=max(0, min(13, f.get('lunch_after', 4, type=int))),
                      period_times=(f.get('period_times') or DEFAULT_TIMES).strip(),
                      default_max_per_day=f.get('max_per_day', 6, type=int), created_by_id=current_user.id)
        db.session.add(plan)
        db.session.commit()
        flash('สร้างแผนตารางสอนแล้ว ต่อไปเพิ่มครูและรายวิชา หรือนำเข้าจาก Excel', 'success')
        return redirect(url_for('tt_plan', pid=plan.id))

    @app.route('/timetable/template.xlsx')
    @login_required
    @role_required('admin')
    def tt_template():
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = 'ตั้งค่า'
        for row in [['หัวข้อ', 'ค่า'], ['ชื่อแผน', 'ตารางสอน ภาคเรียนที่ 1/2570'], ['จำนวนวันต่อสัปดาห์', 5],
                    ['จำนวนคาบต่อวัน', 8], ['พักกลางวันหลังคาบที่', 4], ['เวลาแต่ละคาบ (คั่นด้วย ,)', DEFAULT_TIMES],
                    ['คาบสูงสุดต่อวันของครู (ค่าเริ่มต้น)', 6]]:
            ws.append(row)
        ws = wb.create_sheet('ครู')
        ws.append(['ชื่อครู', 'คาบสูงสุดต่อวัน', 'ห้องประจำ', 'กลุ่มสาระที่สอนได้ (คั่นด้วย ,)'])
        ws.append(['นายตัวอย่าง ใจดี', 6, 'ห้องคณิตศาสตร์1', 'คณิตศาสตร์'])
        ws = wb.create_sheet('รายวิชา')
        ws.append(['ครู', 'รหัสวิชา', 'ชื่อวิชา', 'ชั้น/ห้อง', 'คาบต่อสัปดาห์', 'คาบติดกัน', 'ห้องเรียน'])
        ws.append(['นายตัวอย่าง ใจดี', 'ค21102', 'คณิตศาสตร์ 2', 'ม.1/1', 3, 1, ''])
        ws.append(['นายตัวอย่าง ใจดี', 'ค21102', 'คณิตศาสตร์ 2', 'ม.1/2-3', 3, 1, ''])
        ws = wb.create_sheet('ข้อห้าม')
        ws.append(['ประเภท (ครู/ห้อง/ทั้งหมด)', 'ชื่อครูหรือห้อง', 'วัน (จันทร์-ศุกร์ หรือ ทุกวัน)', 'คาบ (เช่น 1-4, 1,2, ทั้งวัน, เช้า, บ่าย)', 'ระดับ (ห้าม/หลีกเลี่ยง)', 'เหตุผล'])
        ws.append(['ครู', 'นายตัวอย่าง ใจดี', 'จันทร์', 'เช้า', 'ห้าม', 'ประชุมกลุ่มสาระ'])
        ws.append(['ครู', 'นายตัวอย่าง ใจดี', 'ทุกวัน', '1', 'หลีกเลี่ยง', 'ไม่สะดวกคาบเช้า'])
        ws.append(['ห้อง', 'ม.1/1', 'พุธ', '7-8', 'ห้าม', 'กิจกรรมชุมนุม'])
        ws.append(['ทั้งหมด', '', 'ทุกวัน', '8', 'หลีกเลี่ยง', 'หลีกเลี่ยงคาบสุดท้าย'])
        ws = wb.create_sheet('ตารางสอน', 0)
        ws.append(SCHEDULE_HEADERS)
        ws.append(['นายตัวอย่าง ใจดี', 'จันทร์', 1, 'ค21102', 'คณิตศาสตร์ 2', 'ม.1/1', 'ห้องคณิตศาสตร์1', '(ไม่บังคับ: ใส่ตารางที่ใช้อยู่ ถ้าต้องการเริ่มจากตารางเดิม)'])
        for wsx in wb.worksheets:
            for col in 'ABCDEFG':
                wsx.column_dimensions[col].width = 26
        return xlsx_response(wb, 'แม่แบบจัดตารางสอน.xlsx')

    def _rows(ws):
        rows = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r is None or all(v in (None, '') for v in r):
                continue
            rows.append([('' if v is None else (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v).strip())) for v in r])
        return rows

    @app.route('/timetable/import', methods=['POST'])
    @login_required
    @role_required('admin')
    def tt_import():
        from openpyxl import load_workbook
        f = request.files.get('file')
        if not f or not f.filename.lower().endswith(('.xlsx', '.xlsm')):
            flash('กรุณาเลือกไฟล์ Excel (.xlsx)', 'danger')
            return redirect(url_for('tt_home'))
        try:
            wb = load_workbook(f, data_only=True)
        except Exception:
            flash('อ่านไฟล์ Excel ไม่ได้', 'danger')
            return redirect(url_for('tt_home'))
        names = {ws.title.strip(): ws for ws in wb.worksheets}
        # ชีตตารางรายคาบ: 'ตารางสอน' (รูปแบบเดียวกับหน้านำเข้าตารางสอนของระบบ) หรือ 'ตารางเดิม' หรือชีตแรกที่มีหัว ครู/วัน/คาบ
        sched_ws = names.get('ตารางสอน') or names.get('ตารางเดิม')
        if sched_ws is None:
            first = wb.worksheets[0]
            hdr = [str(c.value or '').strip() for c in first[1]]
            if 'วัน' in hdr and 'คาบ' in hdr:
                sched_ws = first
        sched_rows = []
        if sched_ws is not None:
            hdr = [str(c.value or '').strip().lower() for c in sched_ws[1]]

            def col(*keys):
                for k in keys:
                    if k.lower() in hdr:
                        return hdr.index(k.lower())
                return None
            ci = {'t': col('ครู', 'ชื่อครู', 'teacher'), 'd': col('วัน', 'day'), 'p': col('คาบ', 'period'),
                  'code': col('รหัสวิชา', 'code'), 'name': col('ชื่อวิชา', 'subject', 'รายวิชา'),
                  'cls': col('ชั้น/ห้อง', 'ชั้นเรียน', 'classroom', 'ห้อง'), 'room': col('ห้องเรียน', 'room', 'สถานที่')}
            if ci['t'] is None and ci['d'] is None:  # รูปแบบเก่าไม่มีหัวตาราง: ครู|วัน|คาบ|รหัส|ชั้น
                ci = {'t': 0, 'd': 1, 'p': 2, 'code': 3, 'name': None, 'cls': 4, 'room': None}
            for r in _rows(sched_ws):
                g_ = lambda k: (r[ci[k]] if ci[k] is not None and ci[k] < len(r) else '')
                sched_rows.append({'t': g_('t'), 'd': parse_day(g_('d')), 'p': g_('p'), 'code': g_('code'),
                                   'name': g_('name'), 'cls': g_('cls').replace(' ', ''), 'room': g_('room')})
        if 'รายวิชา' not in names and not sched_rows:
            flash('ไม่พบชีต "รายวิชา" หรือชีตตารางสอน (ครู/วัน/คาบ) ในไฟล์ — ดาวน์โหลดแม่แบบเพื่อดูรูปแบบ', 'danger')
            return redirect(url_for('tt_home'))
        cfg = {}
        if 'ตั้งค่า' in names:
            for r in _rows(names['ตั้งค่า']):
                if len(r) >= 2:
                    cfg[r[0]] = r[1]

        def cfg_int(prefix, default):
            for k, v in cfg.items():
                if k.startswith(prefix):
                    try:
                        return int(float(v))
                    except ValueError:
                        return default
            return default
        times = next((v for k, v in cfg.items() if k.startswith('เวลา')), DEFAULT_TIMES)
        plan = TtPlan(school_id=sid(), name=(next((v for k, v in cfg.items() if k.startswith('ชื่อ')), '') or
                                             request.form.get('name') or f.filename.rsplit('.', 1)[0])[:200],
                      days=cfg_int('จำนวนวัน', 5), periods=cfg_int('จำนวนคาบ', 8), lunch_after=cfg_int('พัก', 4),
                      period_times=times, default_max_per_day=cfg_int('คาบสูงสุด', 6), created_by_id=current_user.id)
        db.session.add(plan)
        db.session.flush()
        tmap = {}

        def teacher(name):
            name = (name or '').strip()
            if not name:
                return None
            if name not in tmap:
                u = User.query.filter_by(full_name=name).filter(User.school_id == sid()).first() if hasattr(User, 'school_id') else None
                t = TtTeacher(plan_id=plan.id, name=name[:200], user_id=u.id if u else None)
                db.session.add(t)
                db.session.flush()
                tmap[name] = t
            return tmap[name]
        if 'ครู' in names:
            for r in _rows(names['ครู']):
                t = teacher(r[0])
                if not t:
                    continue
                if len(r) > 1 and r[1].isdigit():
                    t.max_per_day = int(r[1])
                if len(r) > 2:
                    t.room = r[2][:120]
                if len(r) > 3:
                    t.subject_groups = r[3]
        lkey = {}
        n_lessons = 0
        lesson_rows = _rows(names['รายวิชา']) if 'รายวิชา' in names else []
        if not lesson_rows and sched_rows:
            # ไม่มีชีตรายวิชา: สร้างรายวิชาจากตารางรายคาบ (นับคาบ + เดาคาบติดกันจากตาราง)
            agg = {}
            for x in sched_rows:
                if not x['d'] or not str(x['p']).isdigit() or not (x['code'] or x['name']) or not x['cls']:
                    continue
                k = (x['t'], x['code'], x['cls'])
                a = agg.setdefault(k, {'name': x['name'], 'room': x['room'], 'cells': []})
                a['cells'].append((x['d'], int(x['p'])))
            for (tn, code, cls), a in agg.items():
                runs = []
                for d in sorted({c[0] for c in a['cells']}):
                    ps = sorted(p_ for dd, p_ in a['cells'] if dd == d)
                    run = 1
                    for u, v in zip(ps, ps[1:]):
                        if v == u + 1 and u != plan.lunch_after:
                            run += 1
                        else:
                            runs.append(run)
                            run = 1
                    runs.append(run)
                block = max(set(runs), key=runs.count) if runs else 1
                lesson_rows.append([tn, code, a['name'], cls, str(len(a['cells'])), str(block), a['room']])
        for r in lesson_rows:
            r = r + [''] * (7 - len(r))
            t = teacher(r[0])
            if not (r[1] or r[2]) or not r[3]:
                continue
            try:
                periods = int(float(r[4] or 1))
                block = int(float(r[5] or 1))
            except ValueError:
                periods, block = 1, 1
            l = TtLesson(plan_id=plan.id, teacher_id=t.id if t else None, code=r[1][:40], name=r[2][:200],
                         class_label=r[3][:120], groups=','.join(expand_class(r[3])), periods=max(0, periods),
                         block=max(1, min(4, block)), room=(r[6] or (t.room if t else ''))[:120])
            db.session.add(l)
            db.session.flush()
            lkey[(t.name if t else '', r[1], r[3].replace(' ', ''))] = l
            n_lessons += 1
        n_blocks = 0
        if 'ข้อห้าม' in names:
            half_m = list(range(1, plan.lunch_after + 1))
            half_a = list(range(plan.lunch_after + 1, plan.periods + 1))
            for r in _rows(names['ข้อห้าม']):
                r = r + [''] * (6 - len(r))
                scope = 'teacher' if r[0].startswith('ครู') else ('group' if r[0].startswith('ห้อง') or r[0].startswith('ชั้น') else 'all')
                days = list(range(1, plan.days + 1)) if (not r[2] or 'ทุก' in r[2]) else [parse_day(x) for x in re.split(r'[,\s]+', r[2]) if parse_day(x)]
                if 'เช้า' in r[3]:
                    pers = half_m
                elif 'บ่าย' in r[3]:
                    pers = half_a
                else:
                    pers = parse_periods(r[3], plan.periods) or []
                hard = not r[4].startswith('หลีก')
                t = tmap.get(r[1].strip()) if scope == 'teacher' else None
                if scope == 'teacher' and not t:
                    continue
                for d in days:
                    for p_ in pers:
                        db.session.add(TtBlock(plan_id=plan.id, scope=scope, teacher_id=t.id if t else None,
                                               group_name=r[1].replace(' ', '') if scope == 'group' else '',
                                               day=d, period=p_, hard=hard, weight=30, reason=r[5][:200]))
                        n_blocks += 1
        n_slots = 0
        for x in sched_rows:
            l = lkey.get((x['t'].strip(), x['code'], x['cls']))
            if l and x['d'] and str(x['p']).isdigit():
                db.session.add(TtSlot(plan_id=plan.id, lesson_id=l.id, day=x['d'], period=int(x['p']), locked=False))
                n_slots += 1
        db.session.commit()
        flash(f'นำเข้าแล้ว: ครู {len(tmap)} คน, รายวิชา {n_lessons} รายการ, ข้อห้าม {n_blocks} ช่อง, ตารางเดิม {n_slots} คาบ', 'success')
        return redirect(url_for('tt_plan', pid=plan.id))

    # ------------------------------------------------------------------
    # ภาพรวมแผน + จัดอัตโนมัติ
    # ------------------------------------------------------------------
    @app.route('/timetable/<int:pid>')
    @login_required
    @role_required('admin')
    def tt_plan(pid):
        plan = get_plan(pid)
        teachers, lessons, blocks, slots = plan_data(plan)
        issues, missing, bad, soft = analyze(plan, teachers, lessons, blocks, slots)
        per_day, assigned = teacher_loads(plan, teachers, lessons, slots)
        tname = {t.id: t.name for t in teachers}
        lname = {l.id: f'{l.label} {l.class_label}' for l in lessons}
        counts = defaultdict(int)
        for i in issues:
            counts[i['type']] += 1
        spread = [max(v) - min(v) for v in per_day.values() if sum(v)]
        return render_template('timetable/plan.html', plan=plan, teachers=teachers, lessons=lessons, blocks=blocks,
                               slots=slots, issues=issues, missing=missing, counts=counts, tname=tname, lname=lname,
                               per_day=per_day, assigned=assigned, total_periods=sum(l.periods or 0 for l in lessons),
                               locked=sum(1 for s in slots if s.locked),
                               avg_spread=(sum(spread) / len(spread)) if spread else 0, groups=all_groups(lessons))

    @app.route('/timetable/<int:pid>/publish', methods=['GET', 'POST'])
    @login_required
    @role_required('admin')
    def tt_publish(pid):
        plan = get_plan(pid)
        teachers, lessons, blocks, slots = plan_data(plan)
        semester = get_active_semester()
        users = User.query.filter_by(school_id=sid(), role='teacher', is_active=True).order_by(User.full_name).all()
        slotted_lessons = {slot.lesson_id for slot in slots}
        used_teacher_ids = {lesson.teacher_id for lesson in lessons
                            if lesson.teacher_id and lesson.id in slotted_lessons}
        used_teachers = [teacher for teacher in teachers if teacher.id in used_teacher_ids]
        suggestions, suggestion_scores = {}, {}
        for teacher in used_teachers:
            user, score = suggested_user(teacher, users)
            suggestions[teacher.id] = user.id if user else None
            suggestion_scores[teacher.id] = score

        existing = TeachingSchedule.query.filter_by(
            school_id=sid(), semester_id=semester.id
        ).count() if semester else 0
        issues, missing, bad, soft = analyze(plan, teachers, lessons, blocks, slots)
        hard_issues = [issue for issue in issues
                       if issue['type'] in ('teacher_clash', 'group_clash', 'blocked_hard')]

        if request.method == 'POST':
            if not semester:
                flash('ยังไม่มีภาคเรียนปัจจุบัน กรุณาเลือกภาคเรียนก่อนนำตารางไปใช้', 'danger')
                return redirect(url_for('tt_publish', pid=plan.id))
            if not slots:
                flash('แผนนี้ยังไม่มีคาบที่จัดไว้', 'danger')
                return redirect(url_for('tt_publish', pid=plan.id))
            if missing:
                flash('ยังนำไปใช้ไม่ได้ เพราะมีคาบที่วางไม่ครบ', 'danger')
                return redirect(url_for('tt_publish', pid=plan.id))

            mapped = {}
            for teacher in used_teachers:
                user_id = request.form.get(f'teacher_{teacher.id}', type=int)
                user = next((u for u in users if u.id == user_id), None)
                if not user:
                    flash(f'กรุณาเลือกบัญชีครูให้ “{teacher.name}”', 'danger')
                    return redirect(url_for('tt_publish', pid=plan.id))
                mapped[teacher.id] = user

            old_rows = TeachingSchedule.query.filter_by(school_id=sid(), semester_id=semester.id).all()
            old_ids = [row.id for row in old_rows]
            if old_ids and (PeriodLessonLog.query.filter(PeriodLessonLog.schedule_id.in_(old_ids)).first() or
                            Attendance.query.filter(Attendance.schedule_id.in_(old_ids)).first()):
                flash('ตารางเดิมมีข้อมูลเปิดคาบหรือเช็กชื่อแล้ว จึงไม่แทนที่อัตโนมัติ เพื่อป้องกันประวัติสูญหาย', 'danger')
                return redirect(url_for('tt_publish', pid=plan.id))

            lesson_map = {lesson.id: lesson for lesson in lessons}
            time_rows = times_of(plan)
            for old_row in old_rows:
                db.session.delete(old_row)
            db.session.flush()
            created = created_subjects = created_rooms = 0
            for slot in slots:
                lesson = lesson_map.get(slot.lesson_id)
                if not lesson or lesson.teacher_id not in mapped:
                    continue
                user = mapped[lesson.teacher_id]
                subject_name = lesson.label or lesson.name or lesson.code or 'ไม่ระบุรายวิชา'
                subject = Subject.query.filter_by(school_id=sid(), name=subject_name).order_by(Subject.id).first()
                if not subject:
                    subject = Subject(school_id=sid(), name=subject_name[:150], teacher_id=user.id,
                                      credit=1.0, total_periods=max(1, lesson.periods or 1), is_active=True)
                    db.session.add(subject)
                    db.session.flush()
                    created_subjects += 1
                if not TeacherSubject.query.filter_by(teacher_id=user.id, subject_id=subject.id).first():
                    db.session.add(TeacherSubject(teacher_id=user.id, subject_id=subject.id))

                room_label = (lesson.class_label or '').strip() or 'ไม่ระบุห้อง'
                # ม.6/2-3 หมายถึงคาบร่วมของห้อง ม.6/2 และ ม.6/3 ไม่ใช่ห้องใหม่
                # จึงสร้าง schedule ให้ห้องจริงแต่ละห้อง แล้วหน้าตารางค่อยรวมการแสดงผล
                member_room_labels = expand_class(room_label) or [room_label]
                raw_time = time_rows[slot.period - 1] if 0 < slot.period <= len(time_rows) else ''
                match = re.match(r'^\s*(\d{1,2}:\d{2})\s*[-–]\s*(\d{1,2}:\d{2})\s*$', raw_time)
                start_time, end_time = match.groups() if match else ('08:00', '09:00')
                for member_room_label in member_room_labels:
                    room = Classroom.query.filter_by(school_id=sid(), name=member_room_label).order_by(Classroom.id).first()
                    if not room:
                        room = Classroom(school_id=sid(), name=member_room_label[:100], is_active=True)
                        db.session.add(room)
                        db.session.flush()
                        created_rooms += 1
                    if not TeacherClassroom.query.filter_by(teacher_id=user.id, classroom_id=room.id).first():
                        db.session.add(TeacherClassroom(teacher_id=user.id, classroom_id=room.id))
                    if not SubjectClassroom.query.filter_by(subject_id=subject.id, classroom_id=room.id).first():
                        db.session.add(SubjectClassroom(subject_id=subject.id, classroom_id=room.id))
                    db.session.add(TeachingSchedule(
                        school_id=sid(), semester_id=semester.id, teacher_id=user.id,
                        subject_id=subject.id, classroom_id=room.id, weekday=slot.day - 1,
                        period_no=slot.period, start_time=start_time, end_time=end_time,
                        room_name=(lesson.room or room_label)[:100], topic=''
                    ))
                    created += 1

            for teacher in used_teachers:
                teacher.user_id = mapped[teacher.id].id
            db.session.commit()
            details = []
            if created_subjects:
                details.append(f'สร้างรายวิชาใหม่ {created_subjects}')
            if created_rooms:
                details.append(f'สร้างห้องใหม่ {created_rooms}')
            flash(f'นำตารางไปใช้ใน {semester.name} แล้ว {created} คาบ' +
                  (f' · {" · ".join(details)}' if details else ''), 'success')
            first_user = next(iter(mapped.values()), None)
            return redirect(url_for('schedule', teacher_id=first_user.id if first_user else None))

        return render_template('timetable/publish.html', plan=plan, semester=semester, users=users,
                               teachers=used_teachers, suggestions=suggestions,
                               suggestion_scores=suggestion_scores, slots=slots, existing=existing,
                               missing=missing, hard_issues=hard_issues)

    @app.route('/timetable/<int:pid>/settings', methods=['POST'])
    @login_required
    @role_required('admin')
    def tt_settings(pid):
        plan = get_plan(pid)
        f = request.form
        plan.name = (f.get('name') or plan.name).strip()[:200]
        plan.lunch_after = max(0, min(plan.periods, f.get('lunch_after', plan.lunch_after, type=int)))
        plan.period_times = (f.get('period_times') or plan.period_times).strip()
        plan.default_max_per_day = max(0, f.get('max_per_day', plan.default_max_per_day, type=int))
        plan.max_consecutive = max(1, f.get('max_consecutive', plan.max_consecutive or 4, type=int))
        plan.note = (f.get('note') or '').strip()
        db.session.commit()
        flash('บันทึกการตั้งค่าแล้ว', 'success')
        return redirect(url_for('tt_plan', pid=plan.id))

    @app.route('/timetable/<int:pid>/solve', methods=['POST'])
    @login_required
    @role_required('admin')
    def tt_solve(pid):
        plan = get_plan(pid)
        teachers, lessons, blocks, slots = plan_data(plan)
        mode = request.form.get('mode', 'full')
        seconds = max(3, min(60, request.form.get('seconds', 15, type=int)))
        prob = build_problem(plan, teachers, lessons, blocks, slots)
        initial = None
        if mode == 'repair':
            initial = [{'lesson_id': s.lesson_id, 'day': s.day, 'period': s.period} for s in slots if not s.locked]
        res = TS.solve(prob, time_limit=seconds, initial=initial, keep_initial=(mode == 'repair'))
        TtSlot.query.filter_by(plan_id=plan.id, locked=False).delete()
        for p_ in res['placements']:
            if not p_['locked']:
                db.session.add(TtSlot(plan_id=plan.id, lesson_id=p_['lesson_id'], day=p_['day'], period=p_['period'], locked=False))
        plan.last_cost = int(res['cost'])
        plan.last_run_at = datetime.utcnow()
        db.session.commit()
        n_un = sum(u['size'] for u in res['unplaced'])
        word = 'จัดใหม่ทั้งหมด' if mode != 'repair' else 'แก้เฉพาะจุดที่มีปัญหา'
        if n_un:
            flash(f'{word} เสร็จ แต่ยังวางไม่ได้ {n_un} คาบ — ลองลดข้อห้าม/คาบติดกัน หรือเพิ่มเวลาประมวลผล', 'warning')
        else:
            flash(f'{word} เสร็จแล้ว (ประมวลผล {res["iterations"]:,} รอบ)', 'success')
        return redirect(url_for('tt_plan', pid=plan.id))

    @app.route('/timetable/<int:pid>/clear', methods=['POST'])
    @login_required
    @role_required('admin')
    def tt_clear(pid):
        plan = get_plan(pid)
        q = TtSlot.query.filter_by(plan_id=plan.id)
        if request.form.get('keep_locked') == '1':
            q = q.filter_by(locked=False)
        n = q.delete()
        db.session.commit()
        flash(f'ล้างตารางแล้ว {n} คาบ', 'success')
        return redirect(url_for('tt_plan', pid=plan.id))

    @app.route('/timetable/<int:pid>/copy', methods=['POST'])
    @login_required
    @role_required('admin')
    def tt_copy(pid):
        plan = get_plan(pid)
        new = TtPlan(school_id=plan.school_id, name=(plan.name + ' (สำเนา)')[:200], days=plan.days, periods=plan.periods,
                     lunch_after=plan.lunch_after, period_times=plan.period_times,
                     default_max_per_day=plan.default_max_per_day, max_consecutive=plan.max_consecutive,
                     note=plan.note, created_by_id=current_user.id)
        db.session.add(new)
        db.session.flush()
        tm, lm = {}, {}
        for t in TtTeacher.query.filter_by(plan_id=plan.id).all():
            n = TtTeacher(plan_id=new.id, name=t.name, user_id=t.user_id, max_per_day=t.max_per_day,
                          subject_groups=t.subject_groups, room=t.room)
            db.session.add(n); db.session.flush(); tm[t.id] = n.id
        for l in TtLesson.query.filter_by(plan_id=plan.id).all():
            n = TtLesson(plan_id=new.id, teacher_id=tm.get(l.teacher_id), code=l.code, name=l.name,
                         class_label=l.class_label, groups=l.groups, periods=l.periods, block=l.block, room=l.room)
            db.session.add(n); db.session.flush(); lm[l.id] = n.id
        for b in TtBlock.query.filter_by(plan_id=plan.id).all():
            db.session.add(TtBlock(plan_id=new.id, scope=b.scope, teacher_id=tm.get(b.teacher_id), group_name=b.group_name,
                                   day=b.day, period=b.period, hard=b.hard, weight=b.weight, reason=b.reason))
        for s in TtSlot.query.filter_by(plan_id=plan.id).all():
            db.session.add(TtSlot(plan_id=new.id, lesson_id=lm[s.lesson_id], day=s.day, period=s.period, locked=s.locked))
        db.session.commit()
        flash('คัดลอกแผนแล้ว ใช้ทดลองจัดแบบอื่นได้โดยไม่กระทบแผนเดิม', 'success')
        return redirect(url_for('tt_plan', pid=new.id))

    @app.route('/timetable/<int:pid>/delete', methods=['POST'])
    @login_required
    @role_required('admin')
    def tt_delete(pid):
        plan = get_plan(pid)
        for M in (TtSlot, TtBlock, TtLesson, TtTeacher):
            M.query.filter_by(plan_id=plan.id).delete()
        db.session.delete(plan)
        db.session.commit()
        flash('ลบแผนตารางสอนแล้ว', 'success')
        return redirect(url_for('tt_home'))

    # ------------------------------------------------------------------
    # ข้อมูลครู / รายวิชา
    # ------------------------------------------------------------------
    @app.route('/timetable/<int:pid>/data', methods=['GET', 'POST'])
    @login_required
    @role_required('admin')
    def tt_data(pid):
        plan = get_plan(pid)
        if request.method == 'POST':
            f = request.form
            a = f.get('action')
            if a == 'add_teacher':
                name = (f.get('name') or '').strip()
                if name:
                    db.session.add(TtTeacher(plan_id=plan.id, name=name[:200], max_per_day=f.get('max_per_day', type=int),
                                             room=(f.get('room') or '')[:120], subject_groups=f.get('subject_groups') or ''))
                    flash('เพิ่มครูแล้ว', 'success')
            elif a == 'edit_teacher':
                t = db.session.get(TtTeacher, f.get('teacher_id', type=int) or -1)
                if t and t.plan_id == plan.id:
                    t.name = (f.get('name') or t.name).strip()[:200]
                    t.max_per_day = f.get('max_per_day', type=int)
                    t.room = (f.get('room') or '')[:120]
                    t.subject_groups = f.get('subject_groups') or ''
                    flash('บันทึกข้อมูลครูแล้ว', 'success')
            elif a == 'del_teacher':
                t = db.session.get(TtTeacher, f.get('teacher_id', type=int) or -1)
                if t and t.plan_id == plan.id:
                    TtLesson.query.filter_by(teacher_id=t.id).update({'teacher_id': None})
                    TtBlock.query.filter_by(plan_id=plan.id, scope='teacher', teacher_id=t.id).delete()
                    db.session.delete(t)
                    flash('ลบครูแล้ว (รายวิชาของครูคนนี้ยังอยู่ แต่ยังไม่มีผู้สอน)', 'warning')
            elif a == 'add_lesson':
                l = TtLesson(plan_id=plan.id)
                lesson_from_form(plan, l, f)
                if l.class_label and (l.code or l.name):
                    db.session.add(l)
                    flash('เพิ่มรายวิชาแล้ว', 'success')
                else:
                    flash('กรอกรหัส/ชื่อวิชา และชั้น/ห้อง', 'danger')
            elif a == 'edit_lesson':
                l = db.session.get(TtLesson, f.get('lesson_id', type=int) or -1)
                if l and l.plan_id == plan.id:
                    lesson_from_form(plan, l, f)
                    db.session.flush()
                    trim_extra_slots(l)
                    flash('บันทึกรายวิชาแล้ว', 'success')
            elif a == 'del_lesson':
                l = db.session.get(TtLesson, f.get('lesson_id', type=int) or -1)
                if l and l.plan_id == plan.id:
                    TtSlot.query.filter_by(lesson_id=l.id).delete()
                    db.session.delete(l)
                    flash('ลบรายวิชาแล้ว', 'success')
            db.session.commit()
            return redirect(url_for('tt_data', pid=plan.id, _anchor=f.get('anchor') or ''))
        teachers, lessons, blocks, slots = plan_data(plan)
        per_day, assigned = teacher_loads(plan, teachers, lessons, slots)
        placed = defaultdict(int)
        for s in slots:
            placed[s.lesson_id] += 1
        tfilter = request.args.get('teacher_id', type=int)
        shown = [l for l in lessons if not tfilter or l.teacher_id == tfilter]
        shown.sort(key=lambda l: ((l.teacher.name if l.teacher else 'ฮฮ'), l.code, l.class_label))
        return render_template('timetable/data.html', plan=plan, teachers=teachers, lessons=shown, assigned=assigned,
                               placed=placed, tfilter=tfilter, subject_groups=sorted(set(SUBJECT_GROUPS.values())))

    # ------------------------------------------------------------------
    # ข้อห้าม (ครู/ห้อง/ทั้งโรงเรียน)
    # ------------------------------------------------------------------
    @app.route('/timetable/<int:pid>/constraints', methods=['GET', 'POST'])
    @login_required
    @role_required('admin')
    def tt_constraints(pid):
        plan = get_plan(pid)
        teachers, lessons, blocks, slots = plan_data(plan)
        scope = request.values.get('scope', 'teacher')
        if scope not in ('teacher', 'group', 'all'):
            scope = 'teacher'
        ref = request.values.get('ref', '')
        if scope == 'teacher' and not ref and teachers:
            ref = str(teachers[0].id)
        groups = all_groups(lessons)
        if scope == 'group' and not ref and groups:
            ref = groups[0]

        def q_for():
            q = TtBlock.query.filter_by(plan_id=plan.id, scope=scope)
            if scope == 'teacher':
                q = q.filter_by(teacher_id=int(ref) if str(ref).isdigit() else -1)
            elif scope == 'group':
                q = q.filter_by(group_name=ref)
            return q
        if request.method == 'POST':
            f = request.form
            level = f.get('level', 'hard')
            reason = (f.get('reason') or '').strip()[:200]
            cells = f.getlist('cell')  # "d-p"
            existing = {(b.day, b.period): b for b in q_for().all()}
            if f.get('action') == 'preset':
                preset = f.get('preset')
                cells = []
                if preset == 'mon_morning':
                    cells = [f'1-{p}' for p in range(1, plan.lunch_after + 1)]
                    reason = reason or 'ประชุมวันจันทร์เช้า'
                elif preset == 'no_first':
                    cells = [f'{d}-1' for d in range(1, plan.days + 1)]
                    reason = reason or 'ไม่สะดวกคาบเช้า'
                elif preset == 'no_last':
                    cells = [f'{d}-{plan.periods}' for d in range(1, plan.days + 1)]
                    reason = reason or 'หลีกเลี่ยงคาบสุดท้าย'
                elif preset == 'clear':
                    q_for().delete()
                    db.session.commit()
                    flash('ล้างข้อห้ามของรายการนี้แล้ว', 'success')
                    return redirect(url_for('tt_constraints', pid=plan.id, scope=scope, ref=ref))
                for c in cells:
                    d, p_ = [int(x) for x in c.split('-')]
                    b = existing.get((d, p_))
                    if not b:
                        b = TtBlock(plan_id=plan.id, scope=scope, day=d, period=p_,
                                    teacher_id=int(ref) if scope == 'teacher' else None,
                                    group_name=ref if scope == 'group' else '')
                        db.session.add(b)
                    b.hard = (level == 'hard')
                    b.reason = reason
            else:
                wanted = set()
                for c in cells:
                    try:
                        d, p_ = [int(x) for x in c.split('-')]
                    except ValueError:
                        continue
                    wanted.add((d, p_))
                    st = f.get(f'lv_{d}_{p_}', level)
                    b = existing.get((d, p_))
                    if not b:
                        b = TtBlock(plan_id=plan.id, scope=scope, day=d, period=p_,
                                    teacher_id=int(ref) if scope == 'teacher' else None,
                                    group_name=ref if scope == 'group' else '', reason=reason)
                        db.session.add(b)
                    b.hard = (st == 'hard')
                    if reason:
                        b.reason = reason
                for key, b in existing.items():
                    if key not in wanted:
                        db.session.delete(b)
            db.session.commit()
            flash('บันทึกข้อห้ามแล้ว', 'success')
            return redirect(url_for('tt_constraints', pid=plan.id, scope=scope, ref=ref))
        cur = {(b.day, b.period): b for b in q_for().all()}
        summary = defaultdict(lambda: {'hard': 0, 'soft': 0})
        for b in blocks:
            k = ('teacher', b.teacher_id) if b.scope == 'teacher' else (('group', b.group_name) if b.scope == 'group' else ('all', ''))
            summary[k]['hard' if b.hard else 'soft'] += 1
        tname = {t.id: t.name for t in teachers}
        return render_template('timetable/constraints.html', plan=plan, teachers=teachers, groups=groups, scope=scope,
                               ref=ref, cur=cur, times=times_of(plan), summary=dict(summary), tname=tname)

    # ------------------------------------------------------------------
    # ตาราง (รายครู/รายห้อง) + ปรับแก้
    # ------------------------------------------------------------------
    @app.route('/timetable/<int:pid>/grid')
    @login_required
    @role_required('admin')
    def tt_grid(pid):
        plan = get_plan(pid)
        teachers, lessons, blocks, slots = plan_data(plan)
        view = request.args.get('view', 'teacher')
        groups = all_groups(lessons)
        ref = request.args.get('ref', '')
        if view == 'teacher':
            if not ref and teachers:
                ref = str(teachers[0].id)
            sel_lessons = {l.id for l in lessons if str(l.teacher_id) == str(ref)}
            title = next((t.name for t in teachers if str(t.id) == str(ref)), '')
        else:
            view = 'group'
            if not ref and groups:
                ref = groups[0]
            sel_lessons = {l.id for l in lessons if ref in l.group_list}
            title = f'ห้อง {ref}'
        issues, missing, bad, soft = analyze(plan, teachers, lessons, blocks, slots)
        lmap = {l.id: l for l in lessons}
        colors = colors_for(lessons)
        grid = defaultdict(list)
        for s in slots:
            if s.lesson_id in sel_lessons:
                grid[(s.day, s.period)].append(s)
        # ข้อมูลสำหรับ JS ตรวจช่องว่าง
        js = {
            'lessons': {l.id: {'t': l.teacher_id, 'g': l.group_list, 'label': f'{l.label} {l.class_label}', 'p': l.periods}
                        for l in lessons},
            'slots': [{'id': s.id, 'l': s.lesson_id, 'd': s.day, 'p': s.period, 'k': bool(s.locked)} for s in slots],
            'blocks': [{'s': b.scope, 't': b.teacher_id, 'g': b.group_name, 'd': b.day, 'p': b.period, 'h': bool(b.hard)}
                       for b in blocks],
        }
        my_missing = [(l, n) for (l, n) in missing if l.id in sel_lessons]
        blocked_here = {}
        for b in blocks:
            if (view == 'teacher' and b.scope == 'teacher' and str(b.teacher_id) == str(ref)) or \
               (view == 'group' and b.scope == 'group' and b.group_name == ref) or b.scope == 'all':
                prev = blocked_here.get((b.day, b.period))
                blocked_here[(b.day, b.period)] = 'hard' if (b.hard or prev == 'hard') else 'soft'
        per_day, assigned = teacher_loads(plan, teachers, lessons, slots)
        load_row = per_day.get(int(ref)) if view == 'teacher' and str(ref).isdigit() else None
        return render_template('timetable/grid.html', plan=plan, teachers=teachers, groups=groups, view=view, ref=ref,
                               title=title, grid=grid, lmap=lmap, colors=colors, bad=bad, soft=soft, times=times_of(plan),
                               js=json.dumps(js, ensure_ascii=False), my_missing=my_missing, blocked_here=blocked_here,
                               load_row=load_row, n_issues=sum(1 for i in issues if i['type'] in ('teacher_clash', 'group_clash', 'blocked_hard')))

    @app.route('/timetable/<int:pid>/edit', methods=['POST'])
    @login_required
    @role_required('admin')
    def tt_edit(pid):
        plan = get_plan(pid)
        f = request.form
        action = f.get('action')
        d, p_ = f.get('day', type=int), f.get('period', type=int)
        ok = True
        msg = ''
        if action in ('move', 'move_block'):
            s = db.session.get(TtSlot, f.get('slot_id', type=int) or -1)
            if not s or s.plan_id != plan.id or not d or not p_:
                ok, msg = False, 'ไม่พบคาบที่เลือก'
            elif s.locked:
                ok, msg = False, 'คาบนี้ถูกล็อกไว้ ปลดล็อกก่อนย้าย'
            else:
                group = [s]
                if action == 'move_block':
                    # ย้ายคาบติดกันของวิชาเดียวกันในวันเดียวกันไปพร้อมกัน
                    same = {x.period: x for x in TtSlot.query.filter_by(plan_id=plan.id, lesson_id=s.lesson_id, day=s.day).all()}
                    lo = hi = s.period
                    while (lo - 1) in same and not (lo - 1 == plan.lunch_after):
                        lo -= 1
                    while (hi + 1) in same and not (hi == plan.lunch_after):
                        hi += 1
                    group = [same[x] for x in range(lo, hi + 1)]
                    off = p_ - s.period
                    if any(x.locked for x in group):
                        ok, msg = False, 'มีคาบที่ล็อกอยู่ในก้อนนี้'
                    elif not all(1 <= x.period + off <= plan.periods for x in group):
                        ok, msg = False, 'ย้ายทั้งก้อนแล้วเกินจำนวนคาบของวัน'
                    elif len(group) > 1 and plan.lunch_after and                             min(x.period for x in group) + off <= plan.lunch_after < max(x.period for x in group) + off:
                        ok, msg = False, 'ย้ายทั้งก้อนแล้วจะข้ามพักกลางวัน'
                    else:
                        for x in group:
                            x.day, x.period = d, x.period + off
                else:
                    s.day, s.period = d, p_
        elif action == 'place':
            l = db.session.get(TtLesson, f.get('lesson_id', type=int) or -1)
            if not l or l.plan_id != plan.id or not d or not p_:
                ok, msg = False, 'ไม่พบรายวิชา'
            elif TtSlot.query.filter_by(lesson_id=l.id).count() >= (l.periods or 0):
                ok, msg = False, 'วิชานี้วางครบจำนวนคาบแล้ว'
            else:
                db.session.add(TtSlot(plan_id=plan.id, lesson_id=l.id, day=d, period=p_, locked=False))
        elif action in ('lock', 'unplace'):
            s = db.session.get(TtSlot, f.get('slot_id', type=int) or -1)
            if not s or s.plan_id != plan.id:
                ok, msg = False, 'ไม่พบคาบที่เลือก'
            elif action == 'lock':
                s.locked = not s.locked
            else:
                db.session.delete(s)
        elif action == 'lock_view':
            ids = [int(x) for x in f.getlist('slot_ids') if x.isdigit()]
            val = f.get('value') == '1'
            TtSlot.query.filter(TtSlot.plan_id == plan.id, TtSlot.id.in_(ids or [-1])).update({'locked': val}, synchronize_session=False)
        else:
            ok, msg = False, 'คำสั่งไม่ถูกต้อง'
        if ok:
            db.session.commit()
        else:
            db.session.rollback()
        if request.headers.get('X-Requested-With') == 'fetch':
            return jsonify(ok=ok, msg=msg)
        if not ok:
            flash(msg, 'danger')
        return redirect(request.referrer or url_for('tt_grid', pid=plan.id))

    # ------------------------------------------------------------------
    # ภาระงานสอน + ข้อเสนอเกลี่ย
    # ------------------------------------------------------------------
    def suggestions(teachers, lessons, assigned, tolerance=1):
        if not teachers:
            return []
        loads = {t.id: assigned.get(t.id, 0) for t in teachers}
        avg = sum(loads.values()) / len(loads)
        can = {}
        for t in teachers:
            gs = {x.strip() for x in (t.subject_groups or '').split(',') if x.strip()}
            gs.update(subject_group_of(l.code) for l in lessons if l.teacher_id == t.id)
            can[t.id] = gs
        out, used = [], set()
        for _ in range(12):
            best = None
            for l in lessons:
                if not l.teacher_id or l.id in used or l.teacher_id not in loads:
                    continue
                a = l.teacher_id
                if loads[a] <= avg + tolerance:
                    continue
                grp = subject_group_of(l.code)
                for t in teachers:
                    b = t.id
                    if b == a or grp not in can[b] or loads[b] >= avg - tolerance:
                        continue
                    n = l.periods or 0
                    before = (loads[a] - avg) ** 2 + (loads[b] - avg) ** 2
                    after = (loads[a] - n - avg) ** 2 + (loads[b] + n - avg) ** 2
                    gain = before - after
                    if gain > 0 and (best is None or gain > best[0]):
                        best = (gain, l, a, b, n)
            if not best:
                break
            gain, l, a, b, n = best
            loads[a] -= n
            loads[b] += n
            used.add(l.id)
            out.append({'lesson': l, 'from': a, 'to': b, 'periods': n, 'group': subject_group_of(l.code)})
        return out

    @app.route('/timetable/<int:pid>/workload', methods=['GET', 'POST'])
    @login_required
    @role_required('admin')
    def tt_workload(pid):
        plan = get_plan(pid)
        if request.method == 'POST':
            moved = 0
            for item in request.form.getlist('apply'):
                try:
                    lid, to = [int(x) for x in item.split(':')]
                except ValueError:
                    continue
                l = db.session.get(TtLesson, lid)
                t = db.session.get(TtTeacher, to)
                if l and t and l.plan_id == plan.id and t.plan_id == plan.id:
                    l.teacher_id = t.id
                    TtSlot.query.filter_by(lesson_id=l.id).update({'locked': False})
                    moved += 1
            if request.form.get('lesson_id') and request.form.get('teacher_id'):
                l = db.session.get(TtLesson, request.form.get('lesson_id', type=int))
                t = db.session.get(TtTeacher, request.form.get('teacher_id', type=int))
                if l and t and l.plan_id == plan.id and t.plan_id == plan.id:
                    l.teacher_id = t.id
                    moved += 1
            db.session.commit()
            if moved:
                flash(f'โยกรายวิชาแล้ว {moved} รายการ — ตารางเดิมอาจชนกัน กด "แก้เฉพาะจุดที่มีปัญหา" ที่หน้าภาพรวมเพื่อจัดให้ใหม่', 'warning')
            return redirect(url_for('tt_workload', pid=plan.id))
        teachers, lessons, blocks, slots = plan_data(plan)
        per_day, assigned = teacher_loads(plan, teachers, lessons, slots)
        loads = [assigned.get(t.id, 0) for t in teachers]
        avg = (sum(loads) / len(loads)) if loads else 0
        sd = (sum((x - avg) ** 2 for x in loads) / len(loads)) ** 0.5 if loads else 0
        tol = request.args.get('tol', 1, type=int)
        sugg = suggestions(teachers, lessons, assigned, tolerance=tol)
        tname = {t.id: t.name for t in teachers}
        groups_of = defaultdict(set)
        for l in lessons:
            if l.teacher_id:
                groups_of[l.teacher_id].add(subject_group_of(l.code))
        rows = sorted(teachers, key=lambda t: -assigned.get(t.id, 0))
        return render_template('timetable/workload.html', plan=plan, teachers=rows, per_day=per_day, assigned=assigned,
                               avg=avg, sd=sd, sugg=sugg, tname=tname, tol=tol, groups_of=groups_of,
                               maxload=max(loads) if loads else 1, lessons=lessons)

    # ------------------------------------------------------------------
    # ส่งออก / พิมพ์
    # ------------------------------------------------------------------
    def grid_cells(plan, slots, lmap, pred):
        cells = defaultdict(list)
        for s in slots:
            l = lmap.get(s.lesson_id)
            if l and pred(l):
                cells[(s.day, s.period)].append(l)
        return cells

    @app.route('/timetable/<int:pid>/export.xlsx')
    @login_required
    @role_required('admin')
    def tt_export(pid):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
        plan = get_plan(pid)
        teachers, lessons, blocks, slots = plan_data(plan)
        lmap = {l.id: l for l in lessons}
        tname = {t.id: t.name for t in teachers}
        times = times_of(plan)
        thin = Side(style='thin', color='999999')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        wrap = Alignment(wrap_text=True, vertical='center', horizontal='center')
        wb = Workbook()
        ws = wb.active
        ws.title = 'สรุปภาระงาน'
        per_day, assigned = teacher_loads(plan, teachers, lessons, slots)
        ws.append(['ครู'] + [DAY_NAMES[i] for i in range(plan.days)] + ['รวมในตาราง', 'คาบที่ได้รับมอบหมาย'])
        for t in teachers:
            ws.append([t.name] + per_day[t.id] + [sum(per_day[t.id]), assigned.get(t.id, 0)])
        ws.column_dimensions['A'].width = 30

        def sheet(title, header, cells, show):
            w = wb.create_sheet(re.sub(r'[\\/*?:\[\]]', '-', title)[:31])
            w.append([header])
            w['A1'].font = Font(bold=True, size=13)
            w.append(['วัน/คาบ'] + [f'คาบ {i + 1}\n{times[i]}' for i in range(plan.periods)])
            for d in range(1, plan.days + 1):
                row = [DAY_NAMES[d - 1]]
                for p_ in range(1, plan.periods + 1):
                    row.append('\n'.join(show(l) for l in cells.get((d, p_), [])))
                w.append(row)
            for r in w.iter_rows(min_row=2):
                for c in r:
                    c.alignment = wrap
                    c.border = border
            w.column_dimensions['A'].width = 11
            for i in range(plan.periods):
                w.column_dimensions[chr(66 + i)].width = 17
            for i in range(3, plan.days + 3):
                w.row_dimensions[i].height = 48
            if plan.lunch_after and plan.lunch_after < plan.periods:
                col = chr(66 + plan.lunch_after)
                for i in range(2, plan.days + 3):
                    w[f'{col}{i}'].fill = PatternFill('solid', fgColor='F1F5F9')
        for t in teachers:
            cells = grid_cells(plan, slots, lmap, lambda l, t=t: l.teacher_id == t.id)
            sheet(t.name, f'ตารางสอน {t.name} — {plan.name} (รวม {sum(per_day[t.id])} คาบ)', cells,
                  lambda l: f'{l.code} {l.name}\n{l.class_label}' + (f'\n{l.room}' if l.room else ''))
        for gname in all_groups(lessons):
            cells = grid_cells(plan, slots, lmap, lambda l, gname=gname: gname in l.group_list)
            sheet('ห้อง ' + gname, f'ตารางเรียน {gname} — {plan.name}', cells,
                  lambda l: f'{l.code} {l.name}\n{tname.get(l.teacher_id, "")}')
        ws2 = wb.create_sheet('รายการคาบทั้งหมด')
        ws2.append(['ครู', 'วัน', 'คาบ', 'เวลา', 'รหัสวิชา', 'ชื่อวิชา', 'ชั้น/ห้อง', 'ห้องเรียน', 'ล็อก'])
        for s in sorted(slots, key=lambda s: (tname.get(lmap[s.lesson_id].teacher_id, '') if s.lesson_id in lmap else '', s.day, s.period)):
            l = lmap.get(s.lesson_id)
            if not l:
                continue
            ws2.append([tname.get(l.teacher_id, ''), DAY_NAMES[s.day - 1], s.period, times[s.period - 1] if s.period <= len(times) else '',
                        l.code, l.name, l.class_label, l.room, 'ล็อก' if s.locked else ''])
        safe = re.sub(r'[^\w\-ก-๙ ]', '_', plan.name)[:60]
        return xlsx_response(wb, f'{safe}.xlsx')

    @app.route('/timetable/<int:pid>/export-schedule.xlsx')
    @login_required
    @role_required('admin')
    def tt_export_schedule(pid):
        """ไฟล์สำหรับนำเข้าในหน้า "ตารางสอน" ของระบบ (ชีตแรก = ครู/วัน/คาบ/รหัสวิชา/ชื่อวิชา/ชั้น/ห้อง/ห้องเรียน/หัวข้อ)
        และยังนำกลับเข้าระบบจัดตารางได้ เพราะมีชีต ตั้งค่า/ครู/รายวิชา/ข้อห้าม แนบไว้ด้วย"""
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        plan = get_plan(pid)
        teachers, lessons, blocks, slots = plan_data(plan)
        lmap = {l.id: l for l in lessons}
        tname = {t.id: t.name for t in teachers}
        wb = Workbook()
        ws = wb.active
        ws.title = 'ตารางสอน'
        ws.append(SCHEDULE_HEADERS)
        for s in sorted(slots, key=lambda s: (tname.get(lmap[s.lesson_id].teacher_id, '') if s.lesson_id in lmap else '', s.day, s.period)):
            l = lmap.get(s.lesson_id)
            if not l or not l.teacher_id:
                continue
            ws.append([tname.get(l.teacher_id, ''), DAY_NAMES[s.day - 1], s.period, l.code, l.name, l.class_label, l.room or '', ''])
        w = wb.create_sheet('ตั้งค่า')
        for row in [['หัวข้อ', 'ค่า'], ['ชื่อแผน', plan.name], ['จำนวนวันต่อสัปดาห์', plan.days], ['จำนวนคาบต่อวัน', plan.periods],
                    ['พักกลางวันหลังคาบที่', plan.lunch_after], ['เวลาแต่ละคาบ (คั่นด้วย ,)', plan.period_times],
                    ['คาบสูงสุดต่อวันของครู (ค่าเริ่มต้น)', plan.default_max_per_day]]:
            w.append(row)
        w = wb.create_sheet('ครู')
        w.append(['ชื่อครู', 'คาบสูงสุดต่อวัน', 'ห้องประจำ', 'กลุ่มสาระที่สอนได้ (คั่นด้วย ,)'])
        for t in teachers:
            w.append([t.name, t.max_per_day or '', t.room or '', t.subject_groups or ''])
        w = wb.create_sheet('รายวิชา')
        w.append(['ครู', 'รหัสวิชา', 'ชื่อวิชา', 'ชั้น/ห้อง', 'คาบต่อสัปดาห์', 'คาบติดกัน', 'ห้องเรียน'])
        for l in lessons:
            w.append([tname.get(l.teacher_id, ''), l.code, l.name, l.class_label, l.periods, l.block, l.room or ''])
        w = wb.create_sheet('ข้อห้าม')
        w.append(['ประเภท (ครู/ห้อง/ทั้งหมด)', 'ชื่อครูหรือห้อง', 'วัน (จันทร์-ศุกร์ หรือ ทุกวัน)', 'คาบ (เช่น 1-4, 1,2, ทั้งวัน, เช้า, บ่าย)', 'ระดับ (ห้าม/หลีกเลี่ยง)', 'เหตุผล'])
        for b in sorted(blocks, key=lambda b: (b.scope, b.teacher_id or 0, b.group_name or '', b.day, b.period)):
            kind = 'ครู' if b.scope == 'teacher' else ('ห้อง' if b.scope == 'group' else 'ทั้งหมด')
            who = tname.get(b.teacher_id, '') if b.scope == 'teacher' else (b.group_name or '')
            w.append([kind, who, DAY_NAMES[b.day - 1], str(b.period), 'ห้าม' if b.hard else 'หลีกเลี่ยง', b.reason or ''])
        for x in wb.worksheets:
            for c in x[1]:
                c.font = Font(bold=True)
                c.fill = PatternFill('solid', fgColor='DBEAFE')
            for colx, wd in zip('ABCDEFGH', [30, 12, 8, 12, 30, 12, 22, 20]):
                x.column_dimensions[colx].width = wd
        safe = re.sub(r'[^\w\-ก-๙ ]', '_', plan.name)[:50]
        return xlsx_response(wb, f'{safe}_สำหรับนำเข้าตารางสอน.xlsx')

    @app.route('/timetable/<int:pid>/print')
    @login_required
    @role_required('admin')
    def tt_print(pid):
        plan = get_plan(pid)
        teachers, lessons, blocks, slots = plan_data(plan)
        lmap = {l.id: l for l in lessons}
        tname = {t.id: t.name for t in teachers}
        view = request.args.get('view', 'teacher')
        ref = request.args.get('ref', 'all')
        pages = []
        if view == 'teacher':
            for t in teachers:
                if ref != 'all' and str(t.id) != ref:
                    continue
                cells = grid_cells(plan, slots, lmap, lambda l, t=t: l.teacher_id == t.id)
                total = sum(len(v) for v in cells.values())
                pages.append({'title': f'{t.name} จำนวนคาบ {total}',
                              'cells': {k: [{'a': f'{l.code} {l.name}', 'b': l.class_label, 'c': l.room} for l in v] for k, v in cells.items()}})
        else:
            for gname in all_groups(lessons):
                if ref != 'all' and gname != ref:
                    continue
                cells = grid_cells(plan, slots, lmap, lambda l, gname=gname: gname in l.group_list)
                pages.append({'title': f'ห้อง {gname}',
                              'cells': {k: [{'a': f'{l.code} {l.name}', 'b': tname.get(l.teacher_id, ''), 'c': l.room} for l in v] for k, v in cells.items()}})
        return render_template('timetable/print.html', plan=plan, pages=pages, times=times_of(plan))

    return g
